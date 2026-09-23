# ruff: noqa: B008

from datetime import date, datetime, time, timedelta
from typing import Optional
from datetime import timezone
from io import BytesIO

from fastapi import APIRouter, Depends, Form, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from infrastructure.db.deps import get_db
from infrastructure.security.role_guard import require_roles_ids


# ============================================================
# RQ - SELECCION
#
# Bandeja general de requisiciones recibidas por Seleccion.
# - REEMPLAZO y PERSONAL_NUEVO.
# - Consecutivo funcional RQ-GTH-#.
# - Tipificacion/SLA de Seleccion.
# - Fecha efectiva de recibido y dias habiles.
# - Festivos Colombia autogestionados por backend.
# - Candidatos vinculados y cobertura real por CONTRATADO (25).
# - KPI por candidato.
# - Observacion de rechazo de Contratacion.
# - NO USAR CASCADE.
# ============================================================

router = APIRouter(
    prefix="/api/seleccion/rq",
    tags=["Seleccion - RQ"],
)

ROL_SELECCION = 2
require_seleccion_rq = require_roles_ids(ROL_SELECCION)

TZ_COLOMBIA = timezone(timedelta(hours=-5), name="America/Bogota")
ESTADO_CONTRATADO = 25
ESTADO_DESISTE = 27
ESTADO_RECHAZADO = 28


# ============================================================
# HELPERS GENERALES
# ============================================================

def _usuario_actual(current) -> str:
    """
    Obtiene un texto trazable del usuario autenticado.

    Nunca convierte el objeto ORM Usuario completo a texto. Prioriza
    NombreUsuario y luego Usuario; además soporta estructuras anidadas
    que pueda devolver el guard de seguridad.
    """
    if current is None:
        return "SELECCION"

    claves = (
        "NombreUsuario",
        "Usuario",
        "nombre_usuario",
        "username",
        "usuario",
        "CorreoCorporativo",
        "Email",
        "email",
        "sub",
    )

    def _extraer(valor):
        if valor is None:
            return None

        if isinstance(valor, str):
            limpio = valor.strip()
            return limpio[:100] if limpio else None

        if isinstance(valor, dict):
            for clave in claves:
                dato = valor.get(clave)
                if isinstance(dato, str) and dato.strip():
                    return dato.strip()[:100]

                # El contexto autenticado puede traer el modelo ORM Usuario
                # dentro de una clave como "Usuario". En ese caso entramos al
                # objeto y extraemos NombreUsuario/Usuario, sin serializarlo.
                if dato is not None and dato is not valor:
                    encontrado = _extraer(dato)
                    if encontrado:
                        return encontrado

            for clave in ("current_user", "user", "usuario_actual", "usuario_obj"):
                anidado = valor.get(clave)
                if anidado is not None and anidado is not valor:
                    encontrado = _extraer(anidado)
                    if encontrado:
                        return encontrado
            return None

        for atributo in claves:
            dato = getattr(valor, atributo, None)
            if isinstance(dato, str) and dato.strip():
                return dato.strip()[:100]

        for atributo in ("current_user", "user", "usuario_actual", "usuario_obj"):
            anidado = getattr(valor, atributo, None)
            if anidado is not None and anidado is not valor:
                encontrado = _extraer(anidado)
                if encontrado:
                    return encontrado

        return None

    return _extraer(current) or "SELECCION"

def _a_fecha(valor) -> Optional[date]:
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return date.fromisoformat(str(valor)[:10])


def _a_datetime_colombia(valor) -> Optional[datetime]:
    if valor is None:
        return None
    if isinstance(valor, date) and not isinstance(valor, datetime):
        return datetime.combine(valor, time.min, tzinfo=TZ_COLOMBIA)
    if not isinstance(valor, datetime):
        valor = datetime.fromisoformat(str(valor))
    if valor.tzinfo is None:
        return valor.replace(tzinfo=TZ_COLOMBIA)
    return valor.astimezone(TZ_COLOMBIA)


def _siguiente_lunes(fecha: date) -> date:
    """Ley Emiliani: traslada al lunes siguiente cuando no cae lunes."""
    return fecha + timedelta(days=(7 - fecha.weekday()) % 7)


def _festivos_colombia_calculados(anio: int) -> dict[date, str]:
    """
    Genera el calendario base de festivos nacionales de Colombia.

    La tabla FestivoColombia sigue siendo la fuente consultada por el sistema.
    Este calculo sirve para autoprovisionar un año que aun no exista.
    Si una norma futura agrega/cambia un festivo, el catalogo puede ajustarse
    sin cambiar el calculo de dias habiles de las RQ ya registradas.
    """
    # Calculo de Pascua (algoritmo de Meeus/Jones/Butcher).
    # Se implementa aqui para no depender del paquete externo python-dateutil.
    a = anio % 19
    b = anio // 100
    c = anio % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes_pascua = (h + l - 7 * m + 114) // 31
    dia_pascua = ((h + l - 7 * m + 114) % 31) + 1
    pascua = date(anio, mes_pascua, dia_pascua)

    festivos = {
        date(anio, 1, 1): "Año Nuevo",
        _siguiente_lunes(date(anio, 1, 6)): "Día de los Reyes Magos",
        _siguiente_lunes(date(anio, 3, 19)): "Día de San José",
        pascua - timedelta(days=3): "Jueves Santo",
        pascua - timedelta(days=2): "Viernes Santo",
        date(anio, 5, 1): "Día del Trabajo",
        _siguiente_lunes(pascua + timedelta(days=39)): "Ascensión del Señor",
        _siguiente_lunes(pascua + timedelta(days=60)): "Corpus Christi",
        _siguiente_lunes(pascua + timedelta(days=68)): "Sagrado Corazón de Jesús",
        _siguiente_lunes(date(anio, 6, 29)): "San Pedro y San Pablo",
        date(anio, 7, 20): "Día de la Independencia",
        date(anio, 8, 7): "Batalla de Boyacá",
        _siguiente_lunes(date(anio, 8, 15)): "Asunción de la Virgen",
        _siguiente_lunes(date(anio, 10, 12)): "Día de la Raza",
        _siguiente_lunes(date(anio, 11, 1)): "Todos los Santos",
        _siguiente_lunes(date(anio, 11, 11)): "Independencia de Cartagena",
        date(anio, 12, 8): "Inmaculada Concepción",
        date(anio, 12, 25): "Navidad",
    }
    return festivos


def _asegurar_festivos_anio(db: Session, anio: int) -> None:
    """Carga automaticamente el año si no existe ningun festivo para él."""
    existe = db.execute(
        text(
            """
            SELECT 1
            FROM public."FestivoColombia"
            WHERE "Fecha" >= :inicio
              AND "Fecha" < :fin
            LIMIT 1;
            """
        ),
        {"inicio": date(anio, 1, 1), "fin": date(anio + 1, 1, 1)},
    ).first()

    if existe:
        return

    for fecha_festivo, nombre in _festivos_colombia_calculados(anio).items():
        db.execute(
            text(
                """
                INSERT INTO public."FestivoColombia"
                    ("Fecha", "Nombre", "Activo", "UsuarioActualizacion")
                VALUES
                    (:fecha, :nombre, true, 'BACKEND_RQ_SELECCION')
                ON CONFLICT ("Fecha") DO NOTHING;
                """
            ),
            {"fecha": fecha_festivo, "nombre": nombre},
        )
    db.commit()


def _festivos_entre(db: Session, inicio: date, fin: date) -> set[date]:
    if fin < inicio:
        inicio, fin = fin, inicio

    for anio in range(inicio.year, fin.year + 1):
        _asegurar_festivos_anio(db, anio)

    rows = db.execute(
        text(
            """
            SELECT "Fecha"
            FROM public."FestivoColombia"
            WHERE COALESCE("Activo", true) = true
              AND "Fecha" BETWEEN :inicio AND :fin;
            """
        ),
        {"inicio": inicio, "fin": fin},
    ).all()
    return {_a_fecha(row[0]) for row in rows if row[0] is not None}


def _es_habil(fecha: date, festivos: set[date]) -> bool:
    return fecha.weekday() < 5 and fecha not in festivos


def _siguiente_dia_habil(fecha: date, festivos: set[date]) -> date:
    cursor = fecha
    while not _es_habil(cursor, festivos):
        cursor += timedelta(days=1)
    return cursor


def _fecha_recibido_seleccion(db: Session, fecha_envio) -> Optional[date]:
    """
    Regla acordada:
    - Día hábil hasta 12:00:00 -> mismo día.
    - Desde 12:01 -> siguiente día hábil.
    - Sábado/domingo/festivo -> siguiente día hábil.
    """
    dt = _a_datetime_colombia(fecha_envio)
    if dt is None:
        return None

    base = dt.date()
    festivos = _festivos_entre(db, base, base + timedelta(days=14))

    if not _es_habil(base, festivos):
        return _siguiente_dia_habil(base + timedelta(days=1), festivos)

    # La regla de negocio dice "hasta las 12:00" y "desde las 12:01".
    # Por ello, 12:00:xx se conserva en el mismo día; 12:01:00 en adelante avanza.
    if (dt.hour, dt.minute) <= (12, 0):
        return base

    return _siguiente_dia_habil(base + timedelta(days=1), festivos)


def _dias_habiles_inclusivos(db: Session, inicio: Optional[date], fin: Optional[date]) -> int:
    if inicio is None or fin is None or fin < inicio:
        return 0

    festivos = _festivos_entre(db, inicio, fin)
    total = 0
    cursor = inicio
    while cursor <= fin:
        if _es_habil(cursor, festivos):
            total += 1
        cursor += timedelta(days=1)
    return total


def _kpi_candidato(estado_id, dias_gestion: int, dias_maximos) -> Optional[str]:
    if dias_maximos is None:
        return None

    estado_id = int(estado_id) if estado_id is not None else None
    limite = int(dias_maximos)

    if estado_id in {ESTADO_DESISTE, ESTADO_RECHAZADO}:
        return "CANCELADA"

    if estado_id == ESTADO_CONTRATADO:
        return "CUMPLE" if dias_gestion <= limite else "NO CUMPLE"

    return "EN TIEMPO" if dias_gestion <= limite else "NO CUMPLE"


def _estado_bandeja_por_cobertura(cantidad_solicitada: int, candidatos: list[dict]) -> str:
    # Solo cubren la RQ los candidatos con vínculo ACTIVO y contratación válida.
    # Los históricos/inactivos se conservan para trazabilidad, pero no cubren.
    contratados = sum(1 for c in candidatos if c.get("CuentaComoCubierto"))
    activos_vinculados = sum(1 for c in candidatos if c.get("ActivoVinculacion"))

    if cantidad_solicitada > 0 and contratados >= cantidad_solicitada:
        return "CERRADO"
    if activos_vinculados > 0:
        return "EN_PROCESO"
    return "ABIERTO"


# ============================================================
# CONSULTA BASE RQ
# ============================================================

CONSULTA_BASE_RQ_SELECCION = """
    SELECT
        rq."IdRQOperaciones",
        rq."ConsecutivoRQ",
        rq."TipoRQ",
        rq."CantidadSolicitada",
        rq."EnviadoSeleccion",
        rq."FechaEnvioSeleccion",
        rq."IdRetiroLaboral",
        rq."IdPazYSalvo",
        rq."IdRegistroPersonal",
        rq."IdCliente",
        rq."IdUsuarioLider",
        rq."IdPerfilRQ",
        rq."TipoNotificacion",
        rq."FechaRetiro",
        rq."FechaUltimoDiaLaborado",
        rq."Observacion",
        rq."RequiereReemplazo",
        rq."Ciudad",
        rq."IdCiudad",
        rq."IdTipoContrato",
        rq."CargoAprobadoPlanta",
        rq."Turno",
        rq."MotivoVacante",
        rq."IdMotivoVacanteRQ",
        rq."ObservacionCliente",
        rq."FechaRegistro",
        rq."EstadoRQ",
        rq."EnviadoRRLL",
        rq."FechaEnvioRRLL",
        rq."FechaCreacion",
        rq."FechaActualizacion",
        rq."IdTipificacionRQSeleccion",

        trs."Nombre" AS "TipificacionSeleccion",
        trs."DiasMaximosGestion",

        rp."NumeroIdentificacion",
        NULLIF(
            TRIM(COALESCE(rp."Nombres", '') || ' ' || COALESCE(rp."Apellidos", '')),
            ''
        ) AS "NombreCompleto",

        c."Nombre" AS "NombreCliente",
        rq."IdCargo",
        ca."NombreCargo",
        u."NombreUsuario" AS "NombreLider",
        prq."CodigoPerfil",
        prq."DescripcionPerfil",
        prq."Genero" AS "GeneroPerfil",
        prq."NivelEscolaridad" AS "NivelEscolaridadPerfil",
        prq."Observaciones" AS "ObservacionesPerfil",
        ciu."Nombre" AS "NombreCiudadCatalogo",
        COALESCE(NULLIF(TRIM(ciu."Nombre"), ''), NULLIF(TRIM(rq."Ciudad"), '')) AS "NombreCiudadFinal",
        tc."Descripcion" AS "TipoContrato",
        mvrq."Nombre" AS "NombreMotivoVacanteCatalogo",
        COALESCE(NULLIF(TRIM(mvrq."Nombre"), ''), NULLIF(TRIM(rq."MotivoVacante"), '')) AS "NombreMotivoVacanteFinal",
        rl."EstadoCasoRRLL",
        rl."FechaEnvioOperaciones"

    FROM public."RQOperaciones" rq
    LEFT JOIN public."RetiroLaboral" rl
        ON rl."IdRetiroLaboral" = rq."IdRetiroLaboral"
    LEFT JOIN public."RegistroPersonal" rp
        ON rp."IdRegistroPersonal" = rq."IdRegistroPersonal"
    INNER JOIN public."Cliente" c
        ON c."IdCliente" = rq."IdCliente"
    INNER JOIN public."Usuario" u
        ON u."IdUsuario" = rq."IdUsuarioLider"
    LEFT JOIN public."PerfilRQ" prq
        ON prq."IdPerfilRQ" = rq."IdPerfilRQ"
    LEFT JOIN public."Cargo" ca
        ON ca."IdCargo" = rq."IdCargo"
    LEFT JOIN public."Ciudad" ciu
        ON ciu."IdCiudad" = rq."IdCiudad"
    LEFT JOIN public."TipoContrato" tc
        ON tc."IdTipoContrato" = rq."IdTipoContrato"
    LEFT JOIN public."MotivoVacanteRQ" mvrq
        ON mvrq."IdMotivoVacanteRQ" = rq."IdMotivoVacanteRQ"
    LEFT JOIN public."TipificacionRQSeleccion" trs
        ON trs."IdTipificacionRQSeleccion" = rq."IdTipificacionRQSeleccion"
"""


def _consultar_candidatos(db: Session, id_rq_operaciones: int, fecha_recibido: Optional[date], dias_maximos) -> list[dict]:
    rows = db.execute(
        text(
            """
            SELECT
                rc."IdRQCandidato",
                rc."IdRegistroPersonal",
                rc."Activo" AS "ActivoVinculacion",
                rc."FechaVinculacion",
                rc."UsuarioVinculacion",
                rc."Salario",
                rp."NumeroIdentificacion",
                NULLIF(TRIM(COALESCE(rp."Nombres", '') || ' ' || COALESCE(rp."Apellidos", '')), '') AS "NombreCompleto",
                rp."IdEstadoProceso",
                ep."Nombre" AS "EstadoProceso",
                cb."FechaIngreso",
                hent."FechaEntregaContratacion",
                hc."FechaContratado",
                hcan."FechaCancelacion",
                ore."ObservacionesRechazo",
                ore."FechaRechazo"
            FROM public."RQCandidato" rc
            INNER JOIN public."RegistroPersonal" rp
                ON rp."IdRegistroPersonal" = rc."IdRegistroPersonal"
            LEFT JOIN public."EstadoProceso" ep
                ON ep."IdEstadoProceso" = rp."IdEstadoProceso"
            LEFT JOIN LATERAL (
                SELECT cb1."FechaIngreso"
                FROM public."ContratacionBasica" cb1
                WHERE cb1."IdRegistroPersonal" = rp."IdRegistroPersonal"
                ORDER BY cb1."IdContratacionBasica" DESC
                LIMIT 1
            ) cb ON true
            LEFT JOIN LATERAL (
                SELECT (h."FechaMovimiento" AT TIME ZONE 'America/Bogota')::date AS "FechaEntregaContratacion"
                FROM public."HistorialEstadoContratacion" h
                WHERE h."IdRegistroPersonal" = rp."IdRegistroPersonal"
                  AND h."EstadoNuevo" = 24
                ORDER BY h."FechaMovimiento" DESC, h."IdHistorialEstadoContratacion" DESC
                LIMIT 1
            ) hent ON true
            LEFT JOIN LATERAL (
                SELECT (h."FechaMovimiento" AT TIME ZONE 'America/Bogota')::date AS "FechaContratado"
                FROM public."HistorialEstadoContratacion" h
                WHERE h."IdRegistroPersonal" = rp."IdRegistroPersonal"
                  AND h."EstadoNuevo" = 25
                ORDER BY h."FechaMovimiento" DESC
                LIMIT 1
            ) hc ON true
            LEFT JOIN LATERAL (
                SELECT (h."FechaMovimiento" AT TIME ZONE 'America/Bogota')::date AS "FechaCancelacion"
                FROM public."HistorialEstadoContratacion" h
                WHERE h."IdRegistroPersonal" = rp."IdRegistroPersonal"
                  AND h."EstadoNuevo" IN (27, 28)
                ORDER BY h."FechaMovimiento" DESC, h."IdHistorialEstadoContratacion" DESC
                LIMIT 1
            ) hcan ON true
            LEFT JOIN LATERAL (
                SELECT o."ObservacionesRechazo", o."FechaRechazo"
                FROM public."ObsRechazoContratacion" o
                WHERE o."IdRegistroPersonal" = rp."IdRegistroPersonal"
                ORDER BY o."FechaRechazo" DESC, o."IdObsRechazoContratacion" DESC
                LIMIT 1
            ) ore ON true
            WHERE rc."IdRQOperaciones" = :id_rq_operaciones
            ORDER BY rc."FechaVinculacion" ASC, rc."IdRQCandidato" ASC;
            """
        ),
        {"id_rq_operaciones": id_rq_operaciones},
    ).mappings().all()

    hoy = datetime.now(TZ_COLOMBIA).date()
    candidatos = []

    for row in rows:
        estado_id = int(row["IdEstadoProceso"]) if row["IdEstadoProceso"] is not None else None
        fecha_contratado = _a_fecha(row["FechaContratado"])
        fecha_cancelacion = _a_fecha(row["FechaCancelacion"])
        fecha_rechazo = _a_fecha(row["FechaRechazo"])

        if estado_id == ESTADO_CONTRATADO and fecha_contratado:
            fecha_fin_kpi = fecha_contratado
        elif estado_id == ESTADO_DESISTE and fecha_cancelacion:
            # Selección registra el paso real a 27 en HistorialEstadoContratacion.
            # Esa fecha congela los días de gestión del candidato desistido.
            fecha_fin_kpi = fecha_cancelacion
        elif estado_id == ESTADO_RECHAZADO:
            # Contratación conserva FechaRechazo en ObsRechazoContratacion.
            # Si no existiera, usamos como respaldo el historial del estado 28.
            fecha_fin_kpi = fecha_rechazo or fecha_cancelacion or hoy
        else:
            fecha_fin_kpi = hoy

        # Una contratación anterior al recibido efectivo de Selección no puede
        # cubrir esta RQ ni producir artificialmente "0 días / CUMPLE".
        contratacion_temporalmente_valida = bool(
            estado_id == ESTADO_CONTRATADO
            and fecha_contratado is not None
            and fecha_recibido is not None
            and fecha_contratado >= fecha_recibido
        )

        if estado_id == ESTADO_CONTRATADO and not contratacion_temporalmente_valida:
            dias_gestion = None
            kpi = None
        else:
            dias_gestion = _dias_habiles_inclusivos(db, fecha_recibido, fecha_fin_kpi)
            kpi = _kpi_candidato(estado_id, dias_gestion, dias_maximos)

        activo_vinculacion = bool(row["ActivoVinculacion"])
        cuenta_como_cubierto = bool(
            activo_vinculacion and contratacion_temporalmente_valida
        )

        candidatos.append({
            "IdRQCandidato": int(row["IdRQCandidato"]),
            "IdRegistroPersonal": int(row["IdRegistroPersonal"]),
            "ActivoVinculacion": activo_vinculacion,
            "FechaVinculacion": row["FechaVinculacion"],
            "UsuarioVinculacion": row["UsuarioVinculacion"],
            "Salario": float(row["Salario"]) if row["Salario"] is not None else None,
            "NumeroIdentificacion": row["NumeroIdentificacion"],
            "NombreCompleto": row["NombreCompleto"],
            "IdEstadoProceso": estado_id,
            "EstadoProceso": row["EstadoProceso"],
            # Columna L: fecha en que Selección avanza el candidato a Contratación (estado 24).
            "FechaEntregaContratacion": _a_fecha(row["FechaEntregaContratacion"]),
            # Columna T: fecha de ingreso registrada en ContratacionBasica.
            "FechaIngreso": _a_fecha(row["FechaIngreso"]),
            # Se conserva para cobertura/KPI: transición real a CONTRATADO (estado 25).
            "FechaContratado": fecha_contratado,
            "ObservacionContratacion": row["ObservacionesRechazo"],
            "FechaRechazo": row["FechaRechazo"],
            "FechaCancelacion": fecha_cancelacion,
            "DiasGestion": dias_gestion,
            "KPI": kpi,
            "CuentaComoCubierto": cuenta_como_cubierto,
        })

    return candidatos


def _serializar_rq_seleccion(db: Session, row, incluir_candidatos: bool = True) -> dict:
    id_rq = int(row["IdRQOperaciones"])
    cantidad_solicitada = int(row["CantidadSolicitada"] or 1)
    consecutivo = int(row["ConsecutivoRQ"]) if row["ConsecutivoRQ"] is not None else None
    fecha_recibido = _fecha_recibido_seleccion(db, row["FechaEnvioSeleccion"])
    dias_maximos = int(row["DiasMaximosGestion"]) if row["DiasMaximosGestion"] is not None else None

    candidatos = _consultar_candidatos(db, id_rq, fecha_recibido, dias_maximos) if incluir_candidatos else []
    cantidad_contratados = sum(1 for c in candidatos if c["CuentaComoCubierto"])
    cantidad_pendientes = max(cantidad_solicitada - cantidad_contratados, 0)
    estado_bandeja = _estado_bandeja_por_cobertura(cantidad_solicitada, candidatos)

    hoy = datetime.now(TZ_COLOMBIA).date()
    dias_gestion_rq = _dias_habiles_inclusivos(db, fecha_recibido, hoy)

    return {
        "IdRQOperaciones": id_rq,
        "ConsecutivoRQ": consecutivo,
        "CodigoRQ": f"RQ-GTH-{consecutivo}" if consecutivo is not None else None,
        "TipoRQ": row["TipoRQ"],
        "EstadoBandeja": estado_bandeja,
        "EstadoRQ": row["EstadoRQ"],
        "CantidadSolicitada": cantidad_solicitada,
        "CantidadContratados": cantidad_contratados,
        "CantidadPendientes": cantidad_pendientes,
        "CoberturaCompleta": cantidad_contratados >= cantidad_solicitada,
        "EnviadoSeleccion": bool(row["EnviadoSeleccion"]),
        "FechaEnvioSeleccion": row["FechaEnvioSeleccion"],
        "FechaRecibidoSeleccion": fecha_recibido,
        "DiasGestionRQ": dias_gestion_rq,

        "IdTipificacionRQSeleccion": (
            int(row["IdTipificacionRQSeleccion"])
            if row["IdTipificacionRQSeleccion"] is not None else None
        ),
        "TipificacionSeleccion": row["TipificacionSeleccion"],
        "DiasMaximosGestion": dias_maximos,

        "IdRetiroLaboral": int(row["IdRetiroLaboral"]) if row["IdRetiroLaboral"] is not None else None,
        "IdPazYSalvo": int(row["IdPazYSalvo"]) if row["IdPazYSalvo"] is not None else None,
        "IdRegistroPersonal": int(row["IdRegistroPersonal"]) if row["IdRegistroPersonal"] is not None else None,
        "NumeroIdentificacion": row["NumeroIdentificacion"],
        "NombreCompleto": row["NombreCompleto"],
        "IdCliente": int(row["IdCliente"]) if row["IdCliente"] is not None else None,
        "NombreCliente": row["NombreCliente"],
        "IdCargo": int(row["IdCargo"]) if row["IdCargo"] is not None else None,
        "NombreCargo": row["NombreCargo"],
        "IdUsuarioLider": str(row["IdUsuarioLider"]) if row["IdUsuarioLider"] is not None else None,
        "NombreLider": row["NombreLider"],
        "IdPerfilRQ": int(row["IdPerfilRQ"]) if row["IdPerfilRQ"] is not None else None,
        "CodigoPerfil": row["CodigoPerfil"],
        "DescripcionPerfil": row["DescripcionPerfil"],
        "GeneroPerfil": row["GeneroPerfil"],
        "NivelEscolaridadPerfil": row["NivelEscolaridadPerfil"],
        "ObservacionesPerfil": row["ObservacionesPerfil"],
        "TipoNotificacion": row["TipoNotificacion"],
        "FechaRetiro": row["FechaRetiro"],
        "FechaUltimoDiaLaborado": row["FechaUltimoDiaLaborado"],
        "Observacion": row["Observacion"],
        "RequiereReemplazo": bool(row["RequiereReemplazo"]),
        "IdCiudad": int(row["IdCiudad"]) if row["IdCiudad"] is not None else None,
        "Ciudad": row["NombreCiudadFinal"],
        "IdTipoContrato": int(row["IdTipoContrato"]) if row["IdTipoContrato"] is not None else None,
        "TipoContrato": row["TipoContrato"],
        "CargoAprobadoPlanta": row["CargoAprobadoPlanta"],
        "Turno": row["Turno"],
        "IdMotivoVacanteRQ": int(row["IdMotivoVacanteRQ"]) if row["IdMotivoVacanteRQ"] is not None else None,
        "MotivoVacante": row["NombreMotivoVacanteFinal"],
        "ObservacionCliente": row["ObservacionCliente"],
        "FechaRegistro": row["FechaRegistro"],
        "EnviadoRRLL": bool(row["EnviadoRRLL"]),
        "FechaEnvioRRLL": row["FechaEnvioRRLL"],
        "EstadoCasoRRLL": row["EstadoCasoRRLL"],
        "FechaEnvioOperaciones": row["FechaEnvioOperaciones"],
        "FechaCreacion": row["FechaCreacion"],
        "FechaActualizacion": row["FechaActualizacion"],
        "Candidatos": candidatos,
    }


def _obtener_rq_row(db: Session, id_rq_operaciones: int):
    return db.execute(
        text(
            CONSULTA_BASE_RQ_SELECCION
            + """
            WHERE rq."IdRQOperaciones" = :id_rq_operaciones
              AND COALESCE(rq."Activo", true) = true
              AND COALESCE(rq."EnviadoSeleccion", false) = true
              AND rq."FechaEnvioSeleccion" IS NOT NULL
              AND NULLIF(TRIM(COALESCE(rq."TipoRQ", '')), '') IS NOT NULL
            LIMIT 1;
            """
        ),
        {"id_rq_operaciones": id_rq_operaciones},
    ).mappings().first()


def _validar_rq(db: Session, id_rq_operaciones: int):
    row = _obtener_rq_row(db, id_rq_operaciones)
    if not row:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="La RQ no existe o todavía no ha sido enviada formalmente a Selección.",
        )
    return row


# ============================================================
# CATALOGOS
# ============================================================

@router.get("/catalogos/tipificaciones")
def listar_tipificaciones_seleccion(
    db: Session = Depends(get_db),
    current=Depends(require_seleccion_rq),
):
    rows = db.execute(
        text(
            """
            SELECT
                "IdTipificacionRQSeleccion",
                "Nombre",
                "DiasMaximosGestion"
            FROM public."TipificacionRQSeleccion"
            WHERE COALESCE("Activo", true) = true
            ORDER BY "IdTipificacionRQSeleccion";
            """
        )
    ).mappings().all()

    return {"success": True, "data": [dict(r) for r in rows]}


# ============================================================
# LISTADO GENERAL
# ============================================================

@router.get("")
@router.get("/")
def listar_rq_recibidas_seleccion(
    db: Session = Depends(get_db),
    current=Depends(require_seleccion_rq),
):
    rows = db.execute(
        text(
            CONSULTA_BASE_RQ_SELECCION
            + """
            WHERE COALESCE(rq."Activo", true) = true
              AND COALESCE(rq."EnviadoSeleccion", false) = true
              AND rq."FechaEnvioSeleccion" IS NOT NULL
              AND NULLIF(TRIM(COALESCE(rq."TipoRQ", '')), '') IS NOT NULL
            ORDER BY
                rq."FechaEnvioSeleccion" DESC NULLS LAST,
                rq."FechaRegistro" DESC,
                rq."IdRQOperaciones" DESC;
            """
        )
    ).mappings().all()

    data = []

    # Mantiene sincronizado el EstadoRQ técnico con la cobertura real también
    # durante la consulta general. Así, cuando Contratación cambia un candidato
    # a CONTRATADO (25), la RQ no queda visualmente CERRADA mientras en BD
    # permanece EN_PROCESO.
    usuario = _usuario_actual(current)
    cambios_estado = False

    for row in rows:
        item = _serializar_rq_seleccion(db, row)

        if item["CoberturaCompleta"]:
            estado_esperado = "CUBIERTA"
        elif any(c["ActivoVinculacion"] for c in item["Candidatos"]):
            estado_esperado = "EN_PROCESO"
        else:
            estado_esperado = (
                "ENVIADO_RRLL"
                if str(row["TipoRQ"] or "").upper() == "REEMPLAZO"
                else "ENVIADO_SELECCION"
            )

        if str(row["EstadoRQ"] or "") != estado_esperado:
            db.execute(
                text(
                    """
                    UPDATE public."RQOperaciones"
                    SET "EstadoRQ" = :estado,
                        "FechaActualizacion" = CURRENT_TIMESTAMP,
                        "UsuarioActualizacion" = :usuario
                    WHERE "IdRQOperaciones" = :id_rq;
                    """
                ),
                {
                    "estado": estado_esperado,
                    "usuario": usuario,
                    "id_rq": int(row["IdRQOperaciones"]),
                },
            )
            cambios_estado = True
            item["EstadoRQ"] = estado_esperado

        data.append(item)

    if cambios_estado:
        db.commit()

    return {
        "success": True,
        "message": "RQ recibidas por Selección consultadas correctamente.",
        "total": len(data),
        "resumen": {
            "Abiertos": sum(1 for item in data if item["EstadoBandeja"] == "ABIERTO"),
            "EnProceso": sum(1 for item in data if item["EstadoBandeja"] == "EN_PROCESO"),
            "Cerrados": sum(1 for item in data if item["EstadoBandeja"] == "CERRADO"),
        },
        "data": data,
    }


# ============================================================
# ACTUALIZAR TIPIFICACION DE SELECCION
# ============================================================

@router.put("/{id_rq_operaciones}/tipificacion")
def actualizar_tipificacion_rq(
    id_rq_operaciones: int,
    id_tipificacion_rq_seleccion: int = Form(...),
    db: Session = Depends(get_db),
    current=Depends(require_seleccion_rq),
):
    _validar_rq(db, id_rq_operaciones)

    tipificacion = db.execute(
        text(
            """
            SELECT "IdTipificacionRQSeleccion", "Nombre", "DiasMaximosGestion"
            FROM public."TipificacionRQSeleccion"
            WHERE "IdTipificacionRQSeleccion" = :id
              AND COALESCE("Activo", true) = true
            LIMIT 1;
            """
        ),
        {"id": id_tipificacion_rq_seleccion},
    ).mappings().first()

    if not tipificacion:
        raise HTTPException(status_code=400, detail="La tipificación seleccionada no existe o está inactiva.")

    db.execute(
        text(
            """
            UPDATE public."RQOperaciones"
            SET "IdTipificacionRQSeleccion" = :id_tipificacion,
                "FechaActualizacion" = CURRENT_TIMESTAMP,
                "UsuarioActualizacion" = :usuario
            WHERE "IdRQOperaciones" = :id_rq;
            """
        ),
        {
            "id_tipificacion": id_tipificacion_rq_seleccion,
            "usuario": _usuario_actual(current),
            "id_rq": id_rq_operaciones,
        },
    )
    db.commit()

    row = _validar_rq(db, id_rq_operaciones)
    return {
        "success": True,
        "message": "Tipificación de Selección actualizada correctamente.",
        "data": _serializar_rq_seleccion(db, row),
    }


# ============================================================
# VINCULAR CANDIDATO A RQ
# ============================================================

@router.post("/{id_rq_operaciones}/candidatos")
def vincular_candidato_rq(
    id_rq_operaciones: int,
    id_registro_personal: int = Form(...),
    salario: float = Form(...),
    db: Session = Depends(get_db),
    current=Depends(require_seleccion_rq),
):
    row_rq = _validar_rq(db, id_rq_operaciones)
    cantidad_solicitada = int(row_rq["CantidadSolicitada"] or 1)

    if salario <= 0:
        raise HTTPException(status_code=400, detail="El salario debe ser mayor a cero.")

    persona = db.execute(
        text(
            """
            SELECT "IdRegistroPersonal"
            FROM public."RegistroPersonal"
            WHERE "IdRegistroPersonal" = :id
            LIMIT 1;
            """
        ),
        {"id": id_registro_personal},
    ).first()
    if not persona:
        raise HTTPException(status_code=404, detail="El candidato no existe.")

    # Si la persona ya figura como CONTRATADO, una contratación histórica
    # anterior al recibido efectivo de esta RQ no puede asociarse como cobertura.
    estado_candidato = db.execute(
        text(
            """
            SELECT
                rp."IdEstadoProceso",
                hc."FechaContratado"
            FROM public."RegistroPersonal" rp
            LEFT JOIN LATERAL (
                SELECT (h."FechaMovimiento" AT TIME ZONE 'America/Bogota')::date AS "FechaContratado"
                FROM public."HistorialEstadoContratacion" h
                WHERE h."IdRegistroPersonal" = rp."IdRegistroPersonal"
                  AND h."EstadoNuevo" = 25
                ORDER BY h."FechaMovimiento" DESC
                LIMIT 1
            ) hc ON true
            WHERE rp."IdRegistroPersonal" = :id
            LIMIT 1;
            """
        ),
        {"id": id_registro_personal},
    ).mappings().first()

    fecha_recibido_rq = _fecha_recibido_seleccion(db, row_rq["FechaEnvioSeleccion"])
    estado_actual_candidato = (
        int(estado_candidato["IdEstadoProceso"])
        if estado_candidato and estado_candidato["IdEstadoProceso"] is not None
        else None
    )
    fecha_contratado_candidato = (
        _a_fecha(estado_candidato["FechaContratado"])
        if estado_candidato and estado_candidato["FechaContratado"] is not None
        else None
    )

    if (
        estado_actual_candidato == ESTADO_CONTRATADO
        and (
            fecha_contratado_candidato is None
            or fecha_recibido_rq is None
            or fecha_contratado_candidato < fecha_recibido_rq
        )
    ):
        raise HTTPException(
            status_code=409,
            detail=(
                "El candidato ya figura como contratado, pero su fecha de contratación "
                "es anterior al recibido efectivo de esta RQ y no puede contabilizarse "
                "como cobertura."
            ),
        )

    ya_vinculado = db.execute(
        text(
            """
            SELECT "IdRQCandidato", "Activo"
            FROM public."RQCandidato"
            WHERE "IdRQOperaciones" = :id_rq
              AND "IdRegistroPersonal" = :id_persona
            LIMIT 1;
            """
        ),
        {"id_rq": id_rq_operaciones, "id_persona": id_registro_personal},
    ).mappings().first()

    if ya_vinculado and bool(ya_vinculado["Activo"]):
        raise HTTPException(status_code=409, detail="El candidato ya está vinculado a esta RQ.")

    # Los candidatos históricos/inactivos se conservan para trazabilidad y no
    # consumen cupo. La cantidad solicitada limita únicamente las vinculaciones
    # activas simultáneas de la RQ.
    activos_vinculados = db.execute(
        text(
            """
            SELECT COUNT(*)
            FROM public."RQCandidato"
            WHERE "IdRQOperaciones" = :id_rq
              AND COALESCE("Activo", true) = true;
            """
        ),
        {"id_rq": id_rq_operaciones},
    ).scalar() or 0

    if int(activos_vinculados) >= cantidad_solicitada:
        raise HTTPException(
            status_code=409,
            detail="La RQ ya tiene ocupados todos los cupos activos disponibles.",
        )

    # La cobertura definitiva sigue dependiendo exclusivamente de candidatos
    # que hayan llegado realmente a CONTRATADO (estado 25).
    contratados = db.execute(
        text(
            """
            SELECT COUNT(*)
            FROM public."RQCandidato" rc
            INNER JOIN public."RegistroPersonal" rp
                ON rp."IdRegistroPersonal" = rc."IdRegistroPersonal"
            INNER JOIN LATERAL (
                SELECT (h."FechaMovimiento" AT TIME ZONE 'America/Bogota')::date AS "FechaContratado"
                FROM public."HistorialEstadoContratacion" h
                WHERE h."IdRegistroPersonal" = rp."IdRegistroPersonal"
                  AND h."EstadoNuevo" = 25
                ORDER BY h."FechaMovimiento" DESC
                LIMIT 1
            ) hc ON true
            WHERE rc."IdRQOperaciones" = :id_rq
              AND COALESCE(rc."Activo", true) = true
              AND rp."IdEstadoProceso" = 25
              AND hc."FechaContratado" >= :fecha_recibido;
            """
        ),
        {"id_rq": id_rq_operaciones, "fecha_recibido": fecha_recibido_rq},
    ).scalar() or 0

    if int(contratados) >= cantidad_solicitada:
        raise HTTPException(
            status_code=409,
            detail="La RQ ya tiene cubierta la cantidad solicitada.",
        )

    usuario = _usuario_actual(current)

    if ya_vinculado:
        db.execute(
            text(
                """
                UPDATE public."RQCandidato"
                SET "Activo" = true,
                    "Salario" = :salario,
                    "FechaActualizacion" = CURRENT_TIMESTAMP,
                    "UsuarioActualizacion" = :usuario
                WHERE "IdRQCandidato" = :id;
                """
            ),
            {"salario": salario, "usuario": usuario, "id": int(ya_vinculado["IdRQCandidato"])},
        )
    else:
        db.execute(
            text(
                """
                INSERT INTO public."RQCandidato"
                    ("IdRQOperaciones", "IdRegistroPersonal", "Activo", "UsuarioVinculacion", "Salario")
                VALUES
                    (:id_rq, :id_persona, true, :usuario, :salario);
                """
            ),
            {"id_rq": id_rq_operaciones, "id_persona": id_registro_personal, "usuario": usuario, "salario": salario},
        )

    # Guardar el candidato seleccionado lo entrega formalmente a Contratación.
    # Solo registramos historial cuando existe una transición real hacia estado 24.
    if estado_actual_candidato != 24:
        db.execute(
            text(
                """
                UPDATE public."RegistroPersonal"
                SET "IdEstadoProceso" = 24
                WHERE "IdRegistroPersonal" = :id_persona;
                """
            ),
            {"id_persona": id_registro_personal},
        )

        db.execute(
            text(
                """
                INSERT INTO public."HistorialEstadoContratacion"
                    ("IdRegistroPersonal", "EstadoAnterior", "EstadoNuevo",
                     "FechaMovimiento", "UsuarioMovimiento")
                VALUES
                    (:id_persona, :estado_anterior, 24, CURRENT_TIMESTAMP, :usuario);
                """
            ),
            {
                "id_persona": id_registro_personal,
                "estado_anterior": estado_actual_candidato,
                "usuario": usuario,
            },
        )

    # Al existir gestión real, la RQ pasa a EN_PROCESO salvo que ya esté cubierta.
    db.execute(
        text(
            """
            UPDATE public."RQOperaciones"
            SET "EstadoRQ" = CASE
                    WHEN "EstadoRQ" = 'CUBIERTA' THEN "EstadoRQ"
                    ELSE 'EN_PROCESO'
                END,
                "FechaActualizacion" = CURRENT_TIMESTAMP,
                "UsuarioActualizacion" = :usuario
            WHERE "IdRQOperaciones" = :id_rq;
            """
        ),
        {"usuario": usuario, "id_rq": id_rq_operaciones},
    )
    db.commit()

    row = _validar_rq(db, id_rq_operaciones)
    return {
        "success": True,
        "message": "Candidato guardado en la RQ y enviado a Contratación correctamente.",
        "data": _serializar_rq_seleccion(db, row),
    }


# ============================================================
# DESVINCULAR CANDIDATO (TRAZABILIDAD: INACTIVA, NO ELIMINA)
# ============================================================

@router.put("/{id_rq_operaciones}/candidatos/{id_registro_personal}/inactivar")
def inactivar_candidato_rq(
    id_rq_operaciones: int,
    id_registro_personal: int,
    db: Session = Depends(get_db),
    current=Depends(require_seleccion_rq),
):
    _validar_rq(db, id_rq_operaciones)

    result = db.execute(
        text(
            """
            UPDATE public."RQCandidato"
            SET "Activo" = false,
                "FechaActualizacion" = CURRENT_TIMESTAMP,
                "UsuarioActualizacion" = :usuario
            WHERE "IdRQOperaciones" = :id_rq
              AND "IdRegistroPersonal" = :id_persona
              AND COALESCE("Activo", true) = true;
            """
        ),
        {
            "usuario": _usuario_actual(current),
            "id_rq": id_rq_operaciones,
            "id_persona": id_registro_personal,
        },
    )

    if result.rowcount == 0:
        db.rollback()
        raise HTTPException(status_code=404, detail="No existe una vinculación activa de ese candidato con la RQ.")

    db.commit()
    row = _validar_rq(db, id_rq_operaciones)
    return {
        "success": True,
        "message": "Candidato inactivado de la RQ sin eliminar su trazabilidad.",
        "data": _serializar_rq_seleccion(db, row),
    }


# ============================================================
# SINCRONIZAR ESTADO RQ SEGUN COBERTURA REAL
# ============================================================

@router.post("/{id_rq_operaciones}/sincronizar")
def sincronizar_estado_rq(
    id_rq_operaciones: int,
    db: Session = Depends(get_db),
    current=Depends(require_seleccion_rq),
):
    row = _validar_rq(db, id_rq_operaciones)

    # Un candidato que DESISTE (27) o es RECHAZADO (28) deja de ocupar
    # un cupo activo de la RQ, pero su registro se conserva para trazabilidad.
    # La sincronización aplica esta regla automáticamente a los vínculos
    # que todavía estén activos.
    usuario = _usuario_actual(current)
    db.execute(
        text(
            """
            UPDATE public."RQCandidato" rc
            SET "Activo" = false,
                "FechaActualizacion" = CURRENT_TIMESTAMP,
                "UsuarioActualizacion" = :usuario
            FROM public."RegistroPersonal" rp
            WHERE rc."IdRQOperaciones" = :id_rq
              AND rc."IdRegistroPersonal" = rp."IdRegistroPersonal"
              AND COALESCE(rc."Activo", true) = true
              AND rp."IdEstadoProceso" IN (27, 28);
            """
        ),
        {"usuario": usuario, "id_rq": id_rq_operaciones},
    )

    # Volvemos a serializar después de liberar los candidatos cancelados,
    # para calcular bandeja, cupos y cobertura con el estado real.
    data = _serializar_rq_seleccion(db, row)

    if data["CoberturaCompleta"]:
        nuevo_estado = "CUBIERTA"
    elif any(c["ActivoVinculacion"] for c in data["Candidatos"]):
        nuevo_estado = "EN_PROCESO"
    else:
        # Conservamos el estado técnico de envío original. Visualmente será ABIERTO.
        nuevo_estado = "ENVIADO_RRLL" if str(row["TipoRQ"] or "").upper() == "REEMPLAZO" else "ENVIADO_SELECCION"

    db.execute(
        text(
            """
            UPDATE public."RQOperaciones"
            SET "EstadoRQ" = :estado,
                "FechaActualizacion" = CURRENT_TIMESTAMP,
                "UsuarioActualizacion" = :usuario
            WHERE "IdRQOperaciones" = :id_rq;
            """
        ),
        {"estado": nuevo_estado, "usuario": usuario, "id_rq": id_rq_operaciones},
    )
    db.commit()

    row = _validar_rq(db, id_rq_operaciones)
    return {
        "success": True,
        "message": "Estado de la RQ sincronizado con su cobertura real.",
        "data": _serializar_rq_seleccion(db, row),
    }


# ============================================================
# EXPORTAR EXCEL CONSOLIDADO RQ - SELECCION
# ============================================================

def _excel_texto(valor):
    return "" if valor is None else str(valor).strip()


def _excel_fecha(valor):
    if valor is None:
        return None
    if isinstance(valor, datetime) and valor.tzinfo is not None:
        return valor.astimezone(TZ_COLOMBIA).replace(tzinfo=None)
    return valor


def _filas_excel_rq(item: dict) -> list[list]:
    candidatos = item.get("Candidatos") or [None]
    es_reemplazo = str(item.get("TipoRQ") or "").upper() == "REEMPLAZO"
    perfil = _excel_texto(item.get("CodigoPerfil") or item.get("DescripcionPerfil"))
    codigo = _excel_texto(item.get("CodigoRQ")) or f'RQ #{item.get("IdRQOperaciones")}'
    filas = []

    for candidato in candidatos:
        c = candidato or {}
        estado = _excel_texto(c.get("EstadoProceso"))
        if not estado:
            estado = {
                "ABIERTO": "Abierto",
                "EN_PROCESO": "En proceso",
                "CERRADO": "Cerrado",
            }.get(
                _excel_texto(item.get("EstadoBandeja")),
                _excel_texto(item.get("EstadoBandeja")),
            )

        filas.append([
            codigo,                                                     # A
            _excel_texto(item.get("NombreCargo")),                     # B
            perfil,                                                     # C
            _excel_texto(item.get("NombreCliente")),                   # D
            _excel_texto(item.get("MotivoVacante")),                   # E
            _excel_texto(item.get("NombreCompleto")) if es_reemplazo else "",  # F
            _excel_texto(item.get("NumeroIdentificacion")) if es_reemplazo else "",  # G
            _excel_texto(item.get("NombreLider")),                     # H
            _excel_fecha(item.get("FechaRecibidoSeleccion")),          # I
            c.get("DiasGestion") if c.get("DiasGestion") is not None else item.get("DiasGestionRQ"),  # J
            estado,                                                     # K
            _excel_fecha(c.get("FechaEntregaContratacion")),           # L: transición a estado 24
            _excel_texto(c.get("NombreCompleto")),                     # M
            _excel_texto(c.get("NumeroIdentificacion")),               # N
            _excel_texto(item.get("ObservacionCliente") or item.get("Observacion")),  # O
            _excel_fecha(item.get("FechaEnvioSeleccion")),             # P
            _excel_texto(item.get("Ciudad")),                          # Q
            _excel_texto(item.get("TipificacionSeleccion")),           # R
            _excel_texto(c.get("KPI")),                                # S
            _excel_fecha(c.get("FechaIngreso")),                       # T: ContratacionBasica.FechaIngreso
            _excel_texto(c.get("ObservacionContratacion")),            # U
        ])

    return filas


@router.get("/exportar/excel")
def exportar_excel_rq_seleccion(
    db: Session = Depends(get_db),
    current=Depends(require_seleccion_rq),
):
    """
    Exporta todas las RQ recibidas por Selección: abiertas, en proceso y cerradas.

    El archivo conserva la estructura operativa A:U utilizada por Selección,
    incluyendo trazabilidad de candidatos históricos. Los colores del reporte
    son únicamente de presentación; no modifican la lógica de cobertura/KPI.
    """
    rows = db.execute(
        text(
            CONSULTA_BASE_RQ_SELECCION
            + """
            WHERE COALESCE(rq."Activo", true) = true
              AND COALESCE(rq."EnviadoSeleccion", false) = true
              AND rq."FechaEnvioSeleccion" IS NOT NULL
              AND NULLIF(TRIM(COALESCE(rq."TipoRQ", '')), '') IS NOT NULL
            ORDER BY rq."FechaEnvioSeleccion" DESC NULLS LAST,
                     rq."FechaRegistro" DESC,
                     rq."IdRQOperaciones" DESC;
            """
        )
    ).mappings().all()

    items = [_serializar_rq_seleccion(db, row) for row in rows]

    wb = Workbook()
    ws = wb.active
    ws.title = "RQ Selección"
    ws.sheet_view.showGridLines = False

    # ------------------------------------------------------------
    # PALETA CORPORATIVA / ESTILOS BASE
    # ------------------------------------------------------------
    verde_oscuro = "006B4F"
    verde_principal = "008F68"
    verde_claro = "E8F5F0"
    verde_muy_claro = "F4FBF8"
    azul_suave = "EAF2F8"
    gris_fondo = "F7FAFC"
    gris_texto = "475569"
    gris_borde = "D7E0E7"
    blanco = "FFFFFF"

    thin = Side(style="thin", color=gris_borde)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ------------------------------------------------------------
    # TÍTULO Y SUBTÍTULO
    # ------------------------------------------------------------
    ws.merge_cells("A1:U1")
    titulo = ws["A1"]
    titulo.value = "LA PERFECCIÓN  |  REPORTE DE REQUISICIONES - SELECCIÓN"
    titulo.font = Font(bold=True, size=18, color=blanco)
    titulo.fill = PatternFill("solid", fgColor=verde_oscuro)
    titulo.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells("A2:U2")
    subtitulo = ws["A2"]
    subtitulo.value = (
        "Consolidado de RQ recibidas por Selección · "
        f"Generado: {datetime.now(TZ_COLOMBIA).strftime('%d/%m/%Y %I:%M %p')}"
    )
    subtitulo.font = Font(italic=True, size=10, color=gris_texto)
    subtitulo.fill = PatternFill("solid", fgColor=verde_claro)
    subtitulo.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 23

    # ------------------------------------------------------------
    # ENCABEZADOS A:U
    # ------------------------------------------------------------
    headers = [
        "CONSECUTIVO",
        "CARGO",
        "PERFIL",
        "SEDE",
        "MOTIVO DE SOLICITUD",
        "PERSONA A QUIEN SE REEMPLAZA",
        "N° DOC ID DE QUIEN SE REEMPLAZA",
        "SOLICITANTE",
        "F. RECIBIDO SELECCIÓN",
        "DÍAS DE GESTIÓN",
        "ESTADO",
        "F ENT A CONTR",
        "NOMBRE PERSONA CONTRATADA",
        "CEDULA",
        "OBSERVACIÓN OPERACIONES",
        "FECHA CARGUE OPERACIONES",
        "CIUDAD",
        "TIPIFICACIÓN DE CARGOS",
        "KPI",
        "F. DE INGRESO",
        "OBSERVACIONES",
    ]

    header_row = 4
    for i, h in enumerate(headers, 1):
        c = ws.cell(header_row, i, h)
        c.font = Font(bold=True, size=10, color=blanco)
        c.fill = PatternFill("solid", fgColor=verde_principal)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    ws.row_dimensions[header_row].height = 48

    # ------------------------------------------------------------
    # DATOS
    # ------------------------------------------------------------
    r = 5
    for item in items:
        for values in _filas_excel_rq(item):
            for col, value in enumerate(values, 1):
                c = ws.cell(r, col, value)
                c.border = border
                c.font = Font(size=10, color="1F2937")
                c.alignment = Alignment(vertical="top", wrap_text=True)

                # Fechas sin hora.
                if col in (9, 12, 20) and isinstance(value, (date, datetime)):
                    c.number_format = "dd/mm/yyyy"

                # Fecha/hora de cargue desde Operaciones.
                if col == 16 and isinstance(value, datetime):
                    c.number_format = "dd/mm/yyyy hh:mm"

                # Documentos siempre como texto para no perder ceros ni notación.
                if col in (7, 14) and value not in (None, ""):
                    c.number_format = "@"

                # Centrado de columnas operativas cortas.
                if col in (1, 3, 7, 9, 10, 11, 12, 14, 17, 19, 20):
                    c.alignment = Alignment(
                        horizontal="center",
                        vertical="top",
                        wrap_text=True,
                    )

            # Fondo alternado de filas, sin pisar el KPI.
            if r % 2 == 0:
                for col in range(1, 22):
                    if col != 19:
                        ws.cell(r, col).fill = PatternFill("solid", fgColor=gris_fondo)

            # Resalta consecutivo.
            ws.cell(r, 1).font = Font(bold=True, color=verde_oscuro)

            # Resalta días de gestión.
            ws.cell(r, 10).font = Font(bold=True, color="334155")

            # ESTADO (columna K).
            estado = _excel_texto(ws.cell(r, 11).value).upper()
            if estado in ("CONTRATADO", "CERRADO"):
                ws.cell(r, 11).fill = PatternFill("solid", fgColor="DCFCE7")
                ws.cell(r, 11).font = Font(bold=True, color="166534")
            elif estado in ("DESISTE DEL PROCESO", "RECHAZADO", "CANCELADO"):
                ws.cell(r, 11).fill = PatternFill("solid", fgColor="FEE2E2")
                ws.cell(r, 11).font = Font(bold=True, color="991B1B")
            elif estado in ("ABIERTO", "EN PROCESO", "EN_PROCESO"):
                ws.cell(r, 11).fill = PatternFill("solid", fgColor=azul_suave)
                ws.cell(r, 11).font = Font(bold=True, color="1D4ED8")

            # KPI (columna S).
            kpi = _excel_texto(ws.cell(r, 19).value).upper()
            if kpi == "CUMPLE":
                ws.cell(r, 19).fill = PatternFill("solid", fgColor="DCFCE7")
                ws.cell(r, 19).font = Font(bold=True, color="166534")
            elif kpi == "NO CUMPLE":
                ws.cell(r, 19).fill = PatternFill("solid", fgColor="FEE2E2")
                ws.cell(r, 19).font = Font(bold=True, color="991B1B")
            elif kpi == "CANCELADA":
                ws.cell(r, 19).fill = PatternFill("solid", fgColor="E5E7EB")
                ws.cell(r, 19).font = Font(bold=True, color="374151")
            elif kpi == "EN TIEMPO":
                ws.cell(r, 19).fill = PatternFill("solid", fgColor="FEF3C7")
                ws.cell(r, 19).font = Font(bold=True, color="92400E")

            r += 1

    last_row = max(r - 1, header_row)

    # ------------------------------------------------------------
    # TABLA, FILTROS, INMOVILIZACIÓN Y DIMENSIONES
    # ------------------------------------------------------------
    if last_row >= 5:
        tabla = Table(displayName="TablaRQSeleccion", ref=f"A{header_row}:U{last_row}")
        estilo_tabla = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,   # ya controlamos el bandeado para conservar KPI/Estado
            showColumnStripes=False,
        )
        tabla.tableStyleInfo = estilo_tabla
        ws.add_table(tabla)
    else:
        ws.auto_filter.ref = f"A{header_row}:U{header_row}"

    ws.freeze_panes = "A5"

    widths = [
        17, 28, 18, 44, 32, 34, 23,
        23, 21, 17, 25, 19, 35, 19,
        66, 25, 19, 34, 17, 19, 56,
    ]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for rownum in range(5, last_row + 1):
        ws.row_dimensions[rownum].height = 30

    # Fila separadora visual.
    ws.row_dimensions[3].height = 8
    for col in range(1, 22):
        ws.cell(3, col).fill = PatternFill("solid", fgColor=blanco)

    # Impresión más limpia.
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"RQ_Seleccion_{datetime.now(TZ_COLOMBIA).strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )


# ============================================================
# DETALLE - DEBE IR AL FINAL PARA NO CAPTURAR /catalogos/...
# ============================================================

@router.get("/{id_rq_operaciones}")
def obtener_detalle_rq_seleccion(
    id_rq_operaciones: int,
    db: Session = Depends(get_db),
    current=Depends(require_seleccion_rq),
):
    row = _validar_rq(db, id_rq_operaciones)
    return {
        "success": True,
        "message": "Detalle de RQ consultado correctamente.",
        "data": _serializar_rq_seleccion(db, row),
    }
