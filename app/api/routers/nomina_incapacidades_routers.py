from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import hashlib
import secrets

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

from infrastructure.db.deps import get_db
from infrastructure.security.auth_dependencies import get_current_user
from services.email_service import enviar_correo_sin_adjunto


router = APIRouter(
    prefix="/api/nomina-incapacidades",
    tags=["Nómina - Incapacidades"],
)


ROLES_NOMINA_PERMITIDOS = {
    "Administrador",
    "Super Administrador",
    "Nómina",
    "Nomina",
}


NOTIFICACIONES_CORREO_ESTADO = {
    "APROBADA": True,
    "RECHAZADA": True,
    "NEGADO": True,
    "EN PROCESO DE PAGO": True,
    "PAGADO": True,
}


TIPOS_INCAPACIDAD_NOMINA = {
    "INCAPACIDAD_1_2_DIAS": "Incapacidad de 1 y 2 días",
    "ACCIDENTE_TRANSITO": "Accidente de tránsito",
    "ACCIDENTE_TRABAJO": "Accidente de trabajo (ARL)",
    "INCAPACIDAD_3_MAS_DIAS": "Incapacidad de 3 o más días",
    "LICENCIA_MATERNA": "Licencia materna",
    "LICENCIA_PATERNA": "Licencia paterna",
}


DIAS_VIGENCIA_ENLACE_CORRECCION = 15

ORIGENES_FRONTEND_PERMITIDOS = {
    "https://laperfeccion.app",
    "https://qa.laperfeccion.app",
    "http://localhost:5173",
    "http://127.0.0.1:5173",
}

FRONTEND_PRODUCCION = "https://laperfeccion.app"


class RechazoIncapacidadRequest(BaseModel):
    observacion: str = Field(..., min_length=3, max_length=1000)


class RadicarIncapacidadRequest(BaseModel):
    numero_radicado: str = Field(..., min_length=1, max_length=100)
    fecha_radicacion: date


class NegarIncapacidadRequest(BaseModel):
    causal_negacion: str = Field(..., min_length=3, max_length=1000)


class PagarIncapacidadRequest(BaseModel):
    valor_pagado: Decimal = Field(..., gt=0, max_digits=14, decimal_places=2)



class ActualizarTipoIncapacidadRequest(BaseModel):
    tipo_incapacidad: str = Field(..., min_length=1, max_length=60)


class ActualizarDatosIncapacidadRequest(BaseModel):
    tipo_incapacidad: str = Field(..., min_length=1, max_length=60)
    fecha_inicio: date
    dias_incapacidad: int = Field(..., gt=0, le=3650)
    es_prorroga: bool
    concepto_sinergy: str | None = Field(default=None, max_length=10)
    diagnostico: str | None = Field(default=None, max_length=9)


def _obtener_valor_usuario(usuario, *nombres):
    for nombre in nombres:
        if isinstance(usuario, dict) and nombre in usuario:
            return usuario.get(nombre)

        if hasattr(usuario, nombre):
            return getattr(usuario, nombre)

    return None


def _obtener_rol_actual(usuario) -> str:
    rol_directo = _obtener_valor_usuario(
        usuario,
        "role",
        "rol",
        "Rol",
        "NombreRol",
        "nombre_rol",
    )

    if rol_directo:
        return str(rol_directo).strip()

    roles = _obtener_valor_usuario(
        usuario,
        "roles",
        "Roles",
    )

    if isinstance(roles, (list, tuple, set)):
        for rol in roles:
            if isinstance(rol, str):
                valor = rol.strip()
            else:
                valor = (
                    getattr(rol, "Nombre", None)
                    or getattr(rol, "NombreRol", None)
                    or getattr(rol, "nombre", None)
                    or getattr(rol, "nombre_rol", None)
                    or ""
                )

            if valor:
                return str(valor).strip()

    return ""


def _obtener_usuario_gestion(usuario_actual) -> str:
    if isinstance(usuario_actual, dict):
        usuario_obj = usuario_actual.get("usuario")

        if usuario_obj is not None:
            nombre_usuario = getattr(
                usuario_obj,
                "NombreUsuario",
                None,
            )

            if nombre_usuario:
                return str(nombre_usuario).strip()[:150]

        payload = usuario_actual.get("payload")

        if isinstance(payload, dict):
            sub = payload.get("sub")

            if sub:
                return str(sub).strip()[:150]

    valor = _obtener_valor_usuario(
        usuario_actual,
        "NombreUsuario",
        "nombre_usuario",
        "Username",
        "username",
        "Email",
        "email",
        "Correo",
        "correo",
        "Nombre",
        "nombre",
    )

    if valor:
        return str(valor).strip()[:150]

    return "usuario_nomina"


def _validar_acceso_nomina(usuario_actual) -> None:
    rol = _obtener_rol_actual(usuario_actual)

    if rol not in ROLES_NOMINA_PERMITIDOS:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="No tiene permisos para gestionar incapacidades de Nómina.",
        )


def _normalizar_estado(estado) -> str:
    valor = str(estado or "").strip()
    return valor or "REGISTRADA"


def _validar_dias_segun_tipo(
    tipo_incapacidad: str,
    dias_incapacidad: int,
) -> None:
    tipo = str(tipo_incapacidad or "").strip().upper()
    dias = int(dias_incapacidad)

    if tipo == "INCAPACIDAD_1_2_DIAS" and dias not in {1, 2}:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Para el tipo INCAPACIDAD DE 1 Y 2 DÍAS "
                "solo se permite registrar 1 o 2 días."
            ),
        )

    if tipo == "INCAPACIDAD_3_MAS_DIAS" and dias < 3:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Para el tipo INCAPACIDAD DE 3 O MÁS DÍAS "
                "se deben registrar mínimo 3 días."
            ),
        )

    if dias <= 0:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Los días de incapacidad deben ser mayores que cero.",
        )


def _incapacidad_a_dict(fila):
    return {
        "id_incapacidad": fila["IdIncapacidadTrabajador"],
        "id_registro_personal": fila["IdRegistroPersonal"],
        "numero_identificacion": fila["NumeroIdentificacion"],
        "nombres": fila["Nombres"],
        "apellidos": fila["Apellidos"],
        "nombre_completo": (
            f'{fila["Nombres"] or ""} {fila["Apellidos"] or ""}'
        ).strip(),
        "eps": fila["Eps"],
        "tipo_incapacidad": fila["TipoIncapacidad"],
        "descripcion_tipo": fila["DescripcionTipoIncapacidad"],
        "fecha_inicio": fila["FechaInicio"],
        "dias_incapacidad": fila["DiasIncapacidad"],
        "fecha_final": fila["FechaFinal"],
        "es_prorroga": fila["EsProrroga"],
        "concepto_sinergy": fila["ConceptoSinergy"],
        "diagnostico": fila["Diagnostico"],
        "estado": _normalizar_estado(fila["Estado"]),
        "observacion_nomina": fila["ObservacionNomina"],
        "usuario_gestion_nomina": fila["UsuarioGestionNomina"],
        "fecha_gestion_nomina": fila["FechaGestionNomina"],
        "numero_radicado": fila["NumeroRadicado"],
        "fecha_radicacion": fila["FechaRadicacion"],
        "causal_negacion": fila["CausalNegacion"],
        "valor_pagado": fila["ValorPagado"],
        "fecha_creacion": fila["FechaCreacion"],
        "fecha_actualizacion": fila["FechaActualizacion"],
        "total_documentos": int(fila["TotalDocumentos"] or 0),
        "fue_corregida": bool(fila["FueCorregida"]),
        "total_correcciones": int(fila["TotalCorrecciones"] or 0),
        "motivo_ultima_correccion": fila["MotivoUltimaCorreccion"],
        "estado_ultima_correccion": fila["EstadoUltimaCorreccion"],
        "fecha_solicitud_ultima_correccion": fila["FechaSolicitudUltimaCorreccion"],
        "fecha_reenvio_ultima_correccion": fila["FechaReenvioUltimaCorreccion"],
    }


def _documento_a_dict(fila):
    return {
        "id_documento": fila["IdDocumentoIncapacidadTrabajador"],
        "id_incapacidad": fila["IdIncapacidadTrabajador"],
        "tipo_documento": fila["TipoDocumento"],
        "nombre_archivo": fila["NombreArchivo"],
        "formato": fila["Formato"],
        "mime_type": fila["MimeType"],
        "tamano_bytes": fila["TamanoBytes"],
        "fecha_creacion": fila["FechaCreacion"],
    }


def _construir_pdf_consolidado(documentos) -> bytes:
    try:
        try:
            from pypdf import PdfReader, PdfWriter
        except ImportError:
            from PyPDF2 import PdfReader, PdfWriter
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                "No está instalada una librería para unir archivos PDF. "
                "Instala pypdf en el entorno del backend."
            ),
        ) from exc

    try:
        from PIL import Image
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                "No está instalada Pillow, necesaria para convertir "
                "imágenes JPG/PNG a PDF."
            ),
        ) from exc

    writer = PdfWriter()
    buffers_temporales = []

    try:
        for documento in documentos:
            nombre_archivo = str(
                documento["NombreArchivo"] or "documento"
            ).strip()

            contenido = documento["DocumentoCargado"]

            if not contenido:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail=(
                        f'El soporte "{nombre_archivo}" no contiene información.'
                    ),
                )

            contenido_bytes = bytes(contenido)

            formato = str(
                documento["Formato"] or ""
            ).strip().lower().lstrip(".")

            mime_type = str(
                documento["MimeType"] or ""
            ).strip().lower()

            es_pdf = (
                formato == "pdf"
                or mime_type == "application/pdf"
                or nombre_archivo.lower().endswith(".pdf")
            )

            es_imagen = (
                formato in {"jpg", "jpeg", "png"}
                or mime_type in {"image/jpeg", "image/png"}
                or nombre_archivo.lower().endswith(
                    (".jpg", ".jpeg", ".png")
                )
            )

            try:
                if es_pdf:
                    buffer_pdf = BytesIO(contenido_bytes)
                    buffers_temporales.append(buffer_pdf)
                    lector = PdfReader(buffer_pdf)

                elif es_imagen:
                    buffer_imagen = BytesIO(contenido_bytes)
                    buffers_temporales.append(buffer_imagen)

                    imagen = Image.open(buffer_imagen)

                    if imagen.mode in ("RGBA", "LA"):
                        fondo = Image.new(
                            "RGB",
                            imagen.size,
                            (255, 255, 255),
                        )

                        if imagen.mode == "RGBA":
                            fondo.paste(
                                imagen,
                                mask=imagen.getchannel("A"),
                            )
                        else:
                            fondo.paste(
                                imagen.convert("RGBA"),
                                mask=imagen.getchannel("A"),
                            )

                        imagen_pdf = fondo
                    else:
                        imagen_pdf = imagen.convert("RGB")

                    buffer_pdf = BytesIO()
                    buffers_temporales.append(buffer_pdf)

                    imagen_pdf.save(
                        buffer_pdf,
                        format="PDF",
                        resolution=150.0,
                    )
                    buffer_pdf.seek(0)

                    lector = PdfReader(buffer_pdf)

                else:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                        detail=(
                            f'El soporte "{nombre_archivo}" tiene un formato '
                            "que no puede consolidarse en PDF."
                        ),
                    )

                for pagina in lector.pages:
                    writer.add_page(pagina)

            except HTTPException:
                raise
            except Exception as exc:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail=(
                        f'No fue posible procesar el soporte '
                        f'"{nombre_archivo}" para el PDF consolidado.'
                    ),
                ) from exc

        if len(writer.pages) == 0:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No hay páginas disponibles para generar el PDF consolidado.",
            )

        salida = BytesIO()
        writer.write(salida)
        return salida.getvalue()

    finally:
        for buffer in buffers_temporales:
            try:
                buffer.close()
            except Exception:
                pass


def _generar_token_correccion() -> tuple[str, str]:
    token_plano = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(
        token_plano.encode("utf-8")
    ).hexdigest()

    return token_plano, token_hash


def _resolver_origen_frontend(request: Request) -> str:
    origen = str(
        request.headers.get("origin") or ""
    ).strip().rstrip("/")

    if origen in ORIGENES_FRONTEND_PERMITIDOS:
        return origen

    return FRONTEND_PRODUCCION


def _construir_url_correccion(
    request: Request,
    token_plano: str,
) -> str:
    origen = _resolver_origen_frontend(request)

    return (
        f"{origen}/trabajador/incapacidades/"
        f"corregir/{token_plano}"
    )


def _crear_solicitud_correccion(
    db: Session,
    id_incapacidad: int,
    motivo_correccion: str,
    usuario_solicitud: str,
) -> tuple[dict, str]:
    token_plano, token_hash = _generar_token_correccion()

    db.execute(
        text(
            """
            UPDATE public."CorreccionIncapacidadTrabajador"
            SET
                "EstadoCorreccion" = 'CERRADA',
                "Activo" = FALSE
            WHERE
                "IdIncapacidadTrabajador" = :id_incapacidad
                AND "Activo" = TRUE
                AND UPPER(COALESCE("EstadoCorreccion", '')) = 'PENDIENTE'
            """
        ),
        {
            "id_incapacidad": id_incapacidad,
        },
    )

    correccion = db.execute(
        text(
            """
            INSERT INTO public."CorreccionIncapacidadTrabajador"
            (
                "IdIncapacidadTrabajador",
                "MotivoCorreccion",
                "UsuarioSolicitud",
                "FechaSolicitud",
                "EstadoCorreccion",
                "TokenCorreccionHash",
                "FechaExpiracionToken",
                "FechaReenvio",
                "Activo"
            )
            VALUES
            (
                :id_incapacidad,
                :motivo_correccion,
                :usuario_solicitud,
                CURRENT_TIMESTAMP,
                'PENDIENTE',
                :token_hash,
                CURRENT_TIMESTAMP
                    + (:dias_vigencia * INTERVAL '1 day'),
                NULL,
                TRUE
            )
            RETURNING
                "IdCorreccionIncapacidad",
                "IdIncapacidadTrabajador",
                "MotivoCorreccion",
                "UsuarioSolicitud",
                "FechaSolicitud",
                "EstadoCorreccion",
                "FechaExpiracionToken",
                "Activo"
            """
        ),
        {
            "id_incapacidad": id_incapacidad,
            "motivo_correccion": motivo_correccion,
            "usuario_solicitud": usuario_solicitud,
            "token_hash": token_hash,
            "dias_vigencia": DIAS_VIGENCIA_ENLACE_CORRECCION,
        },
    ).mappings().first()

    if not correccion:
        raise RuntimeError(
            "No fue posible crear la solicitud de corrección."
        )

    return dict(correccion), token_plano


def _obtener_datos_correo_incapacidad(
    db: Session,
    id_incapacidad: int,
):
    consulta = text(
        """
        SELECT
            rp."Email",
            rp."Nombres",
            rp."Apellidos",
            i."FechaInicio",
            i."FechaFinal",
            i."DiasIncapacidad",
            i."ValorPagado"
        FROM public."IncapacidadTrabajador" i
        INNER JOIN public."RegistroPersonal" rp
            ON rp."IdRegistroPersonal" = i."IdRegistroPersonal"
        WHERE
            i."IdIncapacidadTrabajador" = :id_incapacidad
            AND i."Activo" = TRUE
        LIMIT 1
        """
    )

    return db.execute(
        consulta,
        {"id_incapacidad": id_incapacidad},
    ).mappings().first()


def _formatear_fecha_correo(fecha) -> str:
    if not fecha:
        return "Sin información"

    try:
        return fecha.strftime("%d/%m/%Y")
    except AttributeError:
        return str(fecha)


def _enviar_notificacion_gestion(
    db: Session,
    id_incapacidad: int,
    estado: str,
    observacion: str | None = None,
    url_correccion: str | None = None,
) -> tuple[bool, str]:
    estado_normalizado = _normalizar_estado(estado).upper()

    if not NOTIFICACIONES_CORREO_ESTADO.get(
        estado_normalizado,
        False,
    ):
        return (
            False,
            "La notificación por correo está desactivada para este estado.",
        )

    datos = _obtener_datos_correo_incapacidad(
        db,
        id_incapacidad,
    )

    if not datos:
        return (
            False,
            "No fue posible obtener los datos del trabajador.",
        )

    destinatario = str(
        datos["Email"] or ""
    ).strip()

    if not destinatario:
        return (
            False,
            "El trabajador no tiene correo registrado.",
        )

    nombre = (
        f'{datos["Nombres"] or ""} {datos["Apellidos"] or ""}'
    ).strip()

    fecha_inicio = _formatear_fecha_correo(
        datos["FechaInicio"]
    )
    fecha_final = _formatear_fecha_correo(
        datos["FechaFinal"]
    )
    dias = datos["DiasIncapacidad"] or 0

    if estado_normalizado == "APROBADA":
        asunto = "Incapacidad aprobada - Aseos La Perfección"
        cuerpo = (
            f"Hola {nombre},\n\n"
            "Te informamos que la incapacidad registrada "
            "fue aprobada por el área de Nómina.\n\n"
            f"Fecha de inicio: {fecha_inicio}\n"
            f"Fecha final: {fecha_final}\n"
            f"Días de incapacidad: {dias}\n\n"
            "Cordialmente,\n"
            "Aseos La Perfección"
        )

    elif estado_normalizado == "RECHAZADA":
        asunto = "Incapacidad rechazada - Aseos La Perfección"

        enlace = str(url_correccion or "").strip()

        cuerpo = (
            f"Hola {nombre},\n\n"
            "Te informamos que la incapacidad registrada "
            "fue rechazada por el área de Nómina y requiere "
            "una corrección.\n\n"
            f"Fecha de inicio: {fecha_inicio}\n"
            f"Fecha final: {fecha_final}\n"
            f"Días de incapacidad: {dias}\n\n"
            "Motivo del rechazo:\n"
            f"{str(observacion or '').strip()}\n\n"
            "Para corregir la información y reenviar la misma "
            "incapacidad para revisión, abre el siguiente enlace "
            "del Portal del Trabajador:\n"
            f"{enlace}\n\n"
            "El enlace es personal, corresponde únicamente a esta "
            "solicitud y tiene una vigencia limitada. No lo compartas "
            "con otras personas.\n\n"
            "Cordialmente,\n"
            "Aseos La Perfección"
        )

    elif estado_normalizado == "NEGADO":
        asunto = "Incapacidad negada - Aseos La Perfección"
        cuerpo = (
            f"Hola {nombre},\n\n"
            "Te informamos que la incapacidad radicada "
            "fue negada por la entidad correspondiente.\n\n"
            f"Fecha de inicio: {fecha_inicio}\n"
            f"Fecha final: {fecha_final}\n"
            f"Días de incapacidad: {dias}\n\n"
            "Causal de negación:\n"
            f"{str(observacion or '').strip()}\n\n"
            "Cordialmente,\n"
            "Aseos La Perfección"
        )

    elif estado_normalizado == "EN PROCESO DE PAGO":
        asunto = (
            "Incapacidad en proceso de pago - "
            "Aseos La Perfección"
        )
        cuerpo = (
            f"Hola {nombre},\n\n"
            "Te informamos que la incapacidad radicada "
            "se encuentra en proceso de pago.\n\n"
            f"Fecha de inicio: {fecha_inicio}\n"
            f"Fecha final: {fecha_final}\n"
            f"Días de incapacidad: {dias}\n\n"
            "Cordialmente,\n"
            "Aseos La Perfección"
        )

    elif estado_normalizado == "PAGADO":
        valor_pagado = datos["ValorPagado"]

        if valor_pagado is None:
            valor_texto = "Sin información"
        else:
            valor_texto = (
                f"${valor_pagado:,.2f}"
                .replace(",", "_")
                .replace(".", ",")
                .replace("_", ".")
            )

        asunto = "Incapacidad pagada - Aseos La Perfección"
        cuerpo = (
            f"Hola {nombre},\n\n"
            "Te informamos que la incapacidad registrada "
            "se encuentra en estado pagado.\n\n"
            f"Fecha de inicio: {fecha_inicio}\n"
            f"Fecha final: {fecha_final}\n"
            f"Días de incapacidad: {dias}\n"
            f"Valor pagado: {valor_texto}\n\n"
            "Cordialmente,\n"
            "Aseos La Perfección"
        )

    else:
        return (
            False,
            "No existe una plantilla de correo para este estado.",
        )

    try:
        enviar_correo_sin_adjunto(
            destinatario=destinatario,
            asunto=asunto,
            cuerpo=cuerpo,
        )

        return (
            True,
            f"Notificación enviada a {destinatario}.",
        )

    except Exception as exc:
        print(
            "No fue posible enviar la notificación "
            f"de incapacidad {id_incapacidad}: {exc}"
        )

        return (
            False,
            "La gestión quedó registrada, pero no fue posible "
            "enviar la notificación por correo.",
        )


def _consultar_estado_incapacidad(
    db: Session,
    id_incapacidad: int,
):
    consulta = text(
        """
        SELECT
            "IdIncapacidadTrabajador",
            "Estado",
            "Activo"
        FROM public."IncapacidadTrabajador"
        WHERE "IdIncapacidadTrabajador" = :id_incapacidad
        LIMIT 1
        """
    )

    return db.execute(
        consulta,
        {"id_incapacidad": id_incapacidad},
    ).mappings().first()


def _resolver_incapacidad_no_actualizada(
    db: Session,
    id_incapacidad: int,
):
    fila = _consultar_estado_incapacidad(
        db,
        id_incapacidad,
    )

    if not fila or fila["Activo"] is not True:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="La incapacidad no existe o no está disponible.",
        )

    estado_actual = _normalizar_estado(fila["Estado"])

    if estado_actual.upper() == "BORRADOR":
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="La incapacidad no existe o no está disponible.",
        )

    raise HTTPException(
        status_code=status.HTTP_409_CONFLICT,
        detail=(
            "La incapacidad ya fue gestionada por Nómina "
            f"y actualmente se encuentra en estado {estado_actual}."
        ),
    )


def _codigo_tipo_incapacidad_excel(
    tipo_incapacidad: str,
    concepto_sinergy: str | None = None,
):
    tipo = str(tipo_incapacidad or "").strip().upper()
    concepto = str(concepto_sinergy or "").strip()

    # Regla principal según el catálogo de conceptos entregado para Sinergy.
    # 1 = Enfermedad general
    # 2 = Maternidad / paternidad
    # 3 = Riesgo laboral
    if concepto in {"1601", "1602", "1702", "1703"}:
        return 1

    if concepto in {"1603", "1604"}:
        return 2

    if concepto in {"1606", "1607", "1704", "1705"}:
        return 3

    # Respaldo para registros antiguos que todavía no tengan concepto.
    if tipo in {
        "INCAPACIDAD_1_2_DIAS",
        "INCAPACIDAD_3_MAS_DIAS",
        "ACCIDENTE_TRANSITO",
    }:
        return 1

    if tipo in {
        "LICENCIA_MATERNA",
        "LICENCIA_PATERNA",
    }:
        return 2

    if tipo == "ACCIDENTE_TRABAJO":
        return 3

    return None


def _normalizar_codigo_diagnostico_excel(diagnostico: str | None) -> str:
    valor = str(diagnostico or "").strip().upper()

    return (
        valor
        .replace(".", "")
        .replace("-", "")
        .replace(" ", "")
    )


def _configurar_excel_incapacidades_aprobadas(ws, filas):
    encabezados = [
        "NÚMERO DE INCAPACIDAD ENTIDAD",
        "FECHA DE EMISIÓN",
        "EMPLEADO",
        "FECHA INICIAL INCAPACIDAD",
        "TOTAL DE DÍAS",
        "TIPO INCAPACIDAD",
        "INDICADOR PRÓRROGA",
        "NÚMERO INCAPACIDAD PRÓRROGA",
        "DIAGNÓSTICO",
        "CONCEPTO",
        "FECHA INICIAL REGISTRO",
        "ESTADO TRÁMITE",
    ]

    leyendas = [
        (
            "CAMPO OBLIGATORIO.\n"
            "- DEBE VENIR EN NÚMERO"
        ),
        (
            "CAMPO OBLIGATORIO.\n"
            "- DEBE VENIR EN FORMATO DD/MM/YYYY."
        ),
        (
            "CÓDIGO DEL EMPLEADO. DEBE EXISTIR EN EL MAESTRO DE PERSONAL. "
            "SI HAY INTEGRACIÓN CON GESTIÓN HUMANA CON MANEJO DE CONSECUTIVO "
            "AUTOMÁTICO PARA CÓDIGOS DE EMPLEADO, EN LUGAR DEL CÓDIGO DEBE "
            "VENIR EL NÚMERO DEL DOCUMENTO.\n\n"
            "NOTA: SI EL SISTEMA ESPECIFICA EN LA IMPORTACIÓN QUE EL CÓDIGO "
            "NO EXISTE, SE DEBE CAMBIAR POR LA CÉDULA DE CIUDADANÍA."
        ),
        (
            "CAMPO OBLIGATORIO.\n"
            "- DEBE VENIR EN FORMATO DD/MM/YYYY."
        ),
        (
            "CAMPO OBLIGATORIO.\n"
            "- DEBE SER EL NÚMERO DE DÍAS EN INCAPACIDAD DEL EMPLEADO."
        ),
        (
            "CAMPO OBLIGATORIO.\n"
            "TIPO DE INCAPACIDAD\n"
            "(1) ENFERMEDAD GENERAL\n"
            "(2) MATERNIDAD\n"
            "(3) RIESGOS PROFESIONALES"
        ),
        (
            "CAMPO OBLIGATORIO.\n"
            "(S) SI ES PRÓRROGA\n"
            "(N) NO ES PRÓRROGA"
        ),
        (
            "CAMPO OPCIONAL.\n"
            'SI EN LA COLUMNA "IND_PRORROGA" VIENE S, SE DEBE INDICAR '
            "LA INCAPACIDAD INICIAL."
        ),
        (
            "CAMPO OBLIGATORIO.\n"
            'DEBE VENIR EL ID DIAGNÓSTICO SEGÚN LA TABLA "DIAGNOSTICO".'
        ),
        (
            "CAMPO OBLIGATORIO.\n"
            'DEBE VENIR EL CÓDIGO DEL CONCEPTO SEGÚN LA TABLA "CONCEPTO".'
        ),
        (
            "CAMPO OPCIONAL. SI EL CONCEPTO MANEJA CANTIDAD. SI NO, ES OPCIONAL.\n"
            "- SI VIENE EN BLANCO, SE ASIGNA LA FECHA INICIAL DE REGRESO.\n"
            "- DEBE VENIR EN FORMATO DD/MM/YYYY.\n"
            "- DEBE ESTAR ENTRE LAS FECHAS INICIAL Y FINAL DEL PERÍODO.\n"
            "- DEBE SER MENOR O IGUAL AL CAMPO FECHA FINAL."
        ),
        (
            "TIPOS DE TRÁMITES:\n\n"
            "01 - EN PROCESO\n"
            "02 - PAGADA\n"
            "03 - DEVUELTA\n"
            "04 - NEGADA\n"
            "05 - TIEMPO NO CUMPLIDO\n"
            "06 - INFORMATIVA\n"
            "07 - INFORMATIVA - PROCESADA\n"
            "08 - PARCIAL EXT\n"
            "09 - TOTAL EXT"
        ),
    ]

    fill_encabezado = PatternFill(
        fill_type="solid",
        fgColor="4A86E8",
    )
    fill_leyenda = PatternFill(
        fill_type="solid",
        fgColor="C9DAF8",
    )
    fill_fila_par = PatternFill(
        fill_type="solid",
        fgColor="F7FAFF",
    )

    fuente_encabezado = Font(
        color="FFFFFF",
        bold=True,
        size=11,
    )
    fuente_leyenda = Font(
        color="000000",
        bold=False,
        size=9,
    )
    fuente_datos = Font(
        color="000000",
        size=10,
    )

    borde_fino = Side(
        style="thin",
        color="7F8C8D",
    )
    borde = Border(
        left=borde_fino,
        right=borde_fino,
        top=borde_fino,
        bottom=borde_fino,
    )

    ws.title = "INCAPACIDADES APROBADAS"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ultima_fila = max(2, len(filas) + 2)
    ws.auto_filter.ref = f"A1:L{ultima_fila}"

    for columna, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(
            row=1,
            column=columna,
            value=encabezado,
        )
        celda.fill = fill_encabezado
        celda.font = fuente_encabezado
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        celda.border = borde

    for columna, leyenda in enumerate(leyendas, start=1):
        celda = ws.cell(
            row=2,
            column=columna,
            value=leyenda.upper(),
        )
        celda.fill = fill_leyenda
        celda.font = fuente_leyenda
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        celda.border = borde

    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 150

    for indice, fila in enumerate(filas, start=3):
        tipo_codigo = _codigo_tipo_incapacidad_excel(
            fila["TipoIncapacidad"],
            fila["ConceptoSinergy"],
        )

        identificacion = str(
            fila["NumeroIdentificacion"] or ""
        ).strip().upper()

        id_diagnostico_sinergy = str(
            fila["IdDiagnosticoSinergy"] or ""
        ).strip().upper() or None

        concepto_sinergy = str(
            fila["ConceptoSinergy"] or ""
        ).strip().upper() or None

        valores = [
            identificacion,
            fila["FechaInicio"],
            identificacion,
            fila["FechaInicio"],
            fila["DiasIncapacidad"],
            tipo_codigo,
            "S" if bool(fila["EsProrroga"]) else "N",
            None,
            id_diagnostico_sinergy,
            concepto_sinergy,
            None,
            "01",
        ]

        for columna, valor in enumerate(valores, start=1):
            if isinstance(valor, str):
                valor = valor.upper()

            celda = ws.cell(
                row=indice,
                column=columna,
                value=valor,
            )
            celda.font = fuente_datos
            celda.border = borde
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            if indice % 2 == 0:
                celda.fill = fill_fila_par

        # Fechas en formato oficial solicitado por Sinergy.
        ws.cell(
            row=indice,
            column=2,
        ).number_format = "dd/mm/yyyy"
        ws.cell(
            row=indice,
            column=4,
        ).number_format = "dd/mm/yyyy"

        # Columnas tratadas expresamente como texto para conservar
        # ceros a la izquierda y evitar conversiones automáticas.
        for columna_texto in (1, 3, 9, 10, 12):
            ws.cell(
                row=indice,
                column=columna_texto,
            ).number_format = "@"

        ws.row_dimensions[indice].height = 24

    anchos = {
        1: 30,
        2: 19,
        3: 24,
        4: 25,
        5: 17,
        6: 21,
        7: 22,
        8: 27,
        9: 18,
        10: 18,
        11: 24,
        12: 24,
    }

    for columna, ancho in anchos.items():
        ws.column_dimensions[
            get_column_letter(columna)
        ].width = ancho

    ws.sheet_view.zoomScale = 80
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.print_title_rows = "1:2"
    ws.print_options.horizontalCentered = True



@router.get("/reporte-excel-aprobadas")
def descargar_excel_incapacidades_aprobadas(
    fecha_inicio: date | None = Query(default=None),
    fecha_fin: date | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    if (
        fecha_inicio is not None
        and fecha_fin is not None
        and fecha_fin < fecha_inicio
    ):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="La fecha fin no puede ser anterior a la fecha inicio.",
        )

    condiciones = [
        'i."Activo" = TRUE',
        "UPPER(COALESCE(i.\"Estado\", '')) IN ('APROBADA', 'PENDIENTE RADICACION')",
    ]
    parametros = {}

    if fecha_inicio is not None:
        condiciones.append('i."FechaInicio" >= :fecha_inicio')
        parametros["fecha_inicio"] = fecha_inicio

    if fecha_fin is not None:
        condiciones.append('i."FechaInicio" <= :fecha_fin')
        parametros["fecha_fin"] = fecha_fin

    consulta_sql = f"""
        SELECT
            i."IdIncapacidadTrabajador",
            rp."NumeroIdentificacion",
            rp."Nombres",
            rp."Apellidos",
            i."TipoIncapacidad",
            i."FechaInicio",
            i."DiasIncapacidad",
            i."EsProrroga",
            i."Diagnostico",
            i."ConceptoSinergy",
            i."Estado",
            diag."IdDiagnostico" AS "IdDiagnosticoSinergy",
            diag."CoincidenciasDiagnostico"
        FROM public."IncapacidadTrabajador" i
        INNER JOIN public."RegistroPersonal" rp
            ON rp."IdRegistroPersonal" = i."IdRegistroPersonal"
        LEFT JOIN LATERAL (
            SELECT
                d."IdDiagnostico",
                COUNT(*) OVER () AS "CoincidenciasDiagnostico"
            FROM public."DiagnosticoSinergy" d
            WHERE
                d."Activo" = TRUE
                AND UPPER(
                    REPLACE(
                        REPLACE(
                            REPLACE(
                                COALESCE(i."Diagnostico", ''),
                                '.',
                                ''
                            ),
                            '-',
                            ''
                        ),
                        ' ',
                        ''
                    )
                ) = UPPER(
                    REPLACE(
                        REPLACE(
                            REPLACE(
                                COALESCE(d."CodigoDiagnostico", ''),
                                '.',
                                ''
                            ),
                            '-',
                            ''
                        ),
                        ' ',
                        ''
                    )
                )
            ORDER BY d."IdDiagnostico" ASC
            LIMIT 1
        ) diag ON TRUE
        WHERE {" AND ".join(condiciones)}
        ORDER BY
            i."FechaInicio" ASC,
            rp."NumeroIdentificacion" ASC,
            i."IdIncapacidadTrabajador" ASC
    """

    filas = db.execute(
        text(consulta_sql),
        parametros,
    ).mappings().all()

    if not filas:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=(
                "No hay incapacidades aprobadas para el rango "
                "de fechas seleccionado."
            ),
        )

    filas_validas = []
    filas_omitidas = []

    for fila in filas:
        id_incapacidad = fila["IdIncapacidadTrabajador"]
        diagnostico = str(fila["Diagnostico"] or "").strip()
        concepto = str(fila["ConceptoSinergy"] or "").strip()
        id_diagnostico = str(
            fila["IdDiagnosticoSinergy"] or ""
        ).strip()

        coincidencias_diagnostico = int(
            fila["CoincidenciasDiagnostico"] or 0
        )

        tipo_excel = _codigo_tipo_incapacidad_excel(
            fila["TipoIncapacidad"],
            concepto,
        )

        motivos = []

        if not diagnostico:
            motivos.append("falta diagnóstico")
        elif not id_diagnostico:
            motivos.append(
                f"diagnóstico {diagnostico} no existe en DiagnosticoSinergy"
            )
        elif coincidencias_diagnostico > 1:
            motivos.append(
                (
                    f"diagnóstico {diagnostico} tiene "
                    f"{coincidencias_diagnostico} coincidencias"
                )
            )

        if not concepto:
            motivos.append("falta concepto Sinergy")

        if tipo_excel is None:
            motivos.append(
                "no fue posible determinar el tipo de incapacidad 1, 2 o 3"
            )

        if motivos:
            filas_omitidas.append(
                {
                    "id_incapacidad": id_incapacidad,
                    "motivos": motivos,
                }
            )
        else:
            filas_validas.append(fila)

    if not filas_validas:
        detalle = "; ".join(
            (
                f'Incapacidad {item["id_incapacidad"]}: '
                + ", ".join(item["motivos"])
            )
            for item in filas_omitidas[:10]
        )

        if len(filas_omitidas) > 10:
            detalle += (
                f"; hay {len(filas_omitidas) - 10} incapacidad(es) "
                "adicional(es) con información pendiente"
            )

        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "No hay incapacidades completas para generar el Excel "
                "de Sinergy en el rango seleccionado. "
                f"{detalle}"
            ),
        )

    wb = Workbook()
    ws = wb.active
    _configurar_excel_incapacidades_aprobadas(
        ws,
        filas_validas,
    )

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)

    sufijo_inicio = (
        fecha_inicio.strftime("%Y%m%d")
        if fecha_inicio
        else "todas"
    )
    sufijo_fin = (
        fecha_fin.strftime("%Y%m%d")
        if fecha_fin
        else "todas"
    )

    nombre_archivo = (
        "Incapacidades_Aprobadas_"
        f"{sufijo_inicio}_{sufijo_fin}.xlsx"
    )

    return StreamingResponse(
        salida,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{nombre_archivo}"'
            ),
            "X-Registros-Exportados": str(len(filas_validas)),
            "X-Registros-Omitidos": str(len(filas_omitidas)),
        },
    )


@router.get("")
def listar_incapacidades_nomina(
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    consulta = text(
        '''
        SELECT
            i."IdIncapacidadTrabajador",
            i."IdRegistroPersonal",
            rp."NumeroIdentificacion",
            rp."Nombres",
            rp."Apellidos",
            COALESCE(te."Descripcion", 'No registrada') AS "Eps",
            i."TipoIncapacidad",
            i."DescripcionTipoIncapacidad",
            i."FechaInicio",
            i."DiasIncapacidad",
            i."FechaFinal",
            i."EsProrroga",
            i."Diagnostico",
            i."ConceptoSinergy",
            i."Estado",
            i."ObservacionNomina",
            i."UsuarioGestionNomina",
            i."FechaGestionNomina",
            i."NumeroRadicado",
            i."FechaRadicacion",
            i."CausalNegacion",
            i."ValorPagado",
            i."FechaCreacion",
            i."FechaActualizacion",
            COUNT(d."IdDocumentoIncapacidadTrabajador") AS "TotalDocumentos",
            COALESCE(corr."FueCorregida", FALSE) AS "FueCorregida",
            COALESCE(corr."TotalCorrecciones", 0) AS "TotalCorrecciones",
            corr."MotivoUltimaCorreccion",
            corr."EstadoUltimaCorreccion",
            corr."FechaSolicitudUltimaCorreccion",
            corr."FechaReenvioUltimaCorreccion"
        FROM public."IncapacidadTrabajador" i
        INNER JOIN public."RegistroPersonal" rp
            ON rp."IdRegistroPersonal" = i."IdRegistroPersonal"
        LEFT JOIN public."TipoEps" te
            ON te."IdTipoEps" = rp."IdTipoEps"
        LEFT JOIN LATERAL (
            SELECT
                EXISTS (
                    SELECT 1
                    FROM public."CorreccionIncapacidadTrabajador" c2
                    WHERE c2."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                      AND UPPER(COALESCE(c2."EstadoCorreccion", '')) = 'REENVIADA'
                ) AS "FueCorregida",
                (
                    SELECT COUNT(*)
                    FROM public."CorreccionIncapacidadTrabajador" c3
                    WHERE c3."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                ) AS "TotalCorrecciones",
                ult."MotivoCorreccion" AS "MotivoUltimaCorreccion",
                ult."EstadoCorreccion" AS "EstadoUltimaCorreccion",
                ult."FechaSolicitud" AS "FechaSolicitudUltimaCorreccion",
                ult."FechaReenvio" AS "FechaReenvioUltimaCorreccion"
            FROM (
                SELECT
                    c1."MotivoCorreccion",
                    c1."EstadoCorreccion",
                    c1."FechaSolicitud",
                    c1."FechaReenvio"
                FROM public."CorreccionIncapacidadTrabajador" c1
                WHERE c1."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                ORDER BY c1."IdCorreccionIncapacidad" DESC
                LIMIT 1
            ) ult
        ) corr ON TRUE
        LEFT JOIN public."DocumentoIncapacidadTrabajador" d
            ON d."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
            AND d."Activo" = TRUE
        WHERE
            i."Activo" = TRUE
            AND COALESCE(i."Estado", '') <> 'BORRADOR'
        GROUP BY
            i."IdIncapacidadTrabajador",
            i."IdRegistroPersonal",
            rp."NumeroIdentificacion",
            rp."Nombres",
            rp."Apellidos",
            te."Descripcion",
            i."TipoIncapacidad",
            i."DescripcionTipoIncapacidad",
            i."FechaInicio",
            i."DiasIncapacidad",
            i."FechaFinal",
            i."EsProrroga",
            i."Diagnostico",
            i."ConceptoSinergy",
            i."Estado",
            i."ObservacionNomina",
            i."UsuarioGestionNomina",
            i."FechaGestionNomina",
            i."NumeroRadicado",
            i."FechaRadicacion",
            i."CausalNegacion",
            i."ValorPagado",
            i."FechaCreacion",
            i."FechaActualizacion",
            corr."FueCorregida",
            corr."TotalCorrecciones",
            corr."MotivoUltimaCorreccion",
            corr."EstadoUltimaCorreccion",
            corr."FechaSolicitudUltimaCorreccion",
            corr."FechaReenvioUltimaCorreccion"
        ORDER BY
            i."FechaCreacion" DESC,
            i."IdIncapacidadTrabajador" DESC
        '''
    )

    filas = db.execute(consulta).mappings().all()

    incapacidades = [
        _incapacidad_a_dict(fila)
        for fila in filas
    ]

    return {
        "success": True,
        "total": len(incapacidades),
        "data": incapacidades,
    }


@router.get("/{id_incapacidad}")
def obtener_detalle_incapacidad_nomina(
    id_incapacidad: int,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    consulta_incapacidad = text(
        '''
        SELECT
            i."IdIncapacidadTrabajador",
            i."IdRegistroPersonal",
            rp."NumeroIdentificacion",
            rp."Nombres",
            rp."Apellidos",
            COALESCE(te."Descripcion", 'No registrada') AS "Eps",
            i."TipoIncapacidad",
            i."DescripcionTipoIncapacidad",
            i."FechaInicio",
            i."DiasIncapacidad",
            i."FechaFinal",
            i."EsProrroga",
            i."Diagnostico",
            i."ConceptoSinergy",
            i."Estado",
            i."ObservacionNomina",
            i."UsuarioGestionNomina",
            i."FechaGestionNomina",
            i."NumeroRadicado",
            i."FechaRadicacion",
            i."CausalNegacion",
            i."ValorPagado",
            i."FechaCreacion",
            i."FechaActualizacion",
            (
                SELECT COUNT(*)
                FROM public."DocumentoIncapacidadTrabajador" d
                WHERE
                    d."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                    AND d."Activo" = TRUE
            ) AS "TotalDocumentos",
            EXISTS (
                SELECT 1
                FROM public."CorreccionIncapacidadTrabajador" c2
                WHERE c2."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                  AND UPPER(COALESCE(c2."EstadoCorreccion", '')) = 'REENVIADA'
            ) AS "FueCorregida",
            (
                SELECT COUNT(*)
                FROM public."CorreccionIncapacidadTrabajador" c3
                WHERE c3."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
            ) AS "TotalCorrecciones",
            (
                SELECT c1."MotivoCorreccion"
                FROM public."CorreccionIncapacidadTrabajador" c1
                WHERE c1."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                ORDER BY c1."IdCorreccionIncapacidad" DESC
                LIMIT 1
            ) AS "MotivoUltimaCorreccion",
            (
                SELECT c1."EstadoCorreccion"
                FROM public."CorreccionIncapacidadTrabajador" c1
                WHERE c1."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                ORDER BY c1."IdCorreccionIncapacidad" DESC
                LIMIT 1
            ) AS "EstadoUltimaCorreccion",
            (
                SELECT c1."FechaSolicitud"
                FROM public."CorreccionIncapacidadTrabajador" c1
                WHERE c1."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                ORDER BY c1."IdCorreccionIncapacidad" DESC
                LIMIT 1
            ) AS "FechaSolicitudUltimaCorreccion",
            (
                SELECT c1."FechaReenvio"
                FROM public."CorreccionIncapacidadTrabajador" c1
                WHERE c1."IdIncapacidadTrabajador" = i."IdIncapacidadTrabajador"
                ORDER BY c1."IdCorreccionIncapacidad" DESC
                LIMIT 1
            ) AS "FechaReenvioUltimaCorreccion"
        FROM public."IncapacidadTrabajador" i
        INNER JOIN public."RegistroPersonal" rp
            ON rp."IdRegistroPersonal" = i."IdRegistroPersonal"
        LEFT JOIN public."TipoEps" te
            ON te."IdTipoEps" = rp."IdTipoEps"
        WHERE
            i."IdIncapacidadTrabajador" = :id_incapacidad
            AND i."Activo" = TRUE
            AND COALESCE(i."Estado", '') <> 'BORRADOR'
        LIMIT 1
        '''
    )

    fila = db.execute(
        consulta_incapacidad,
        {"id_incapacidad": id_incapacidad},
    ).mappings().first()

    if not fila:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="La incapacidad no existe o no está disponible.",
        )

    consulta_documentos = text(
        '''
        SELECT
            d."IdDocumentoIncapacidadTrabajador",
            d."IdIncapacidadTrabajador",
            d."TipoDocumento",
            d."NombreArchivo",
            d."Formato",
            d."MimeType",
            d."TamanoBytes",
            d."FechaCreacion"
        FROM public."DocumentoIncapacidadTrabajador" d
        WHERE
            d."IdIncapacidadTrabajador" = :id_incapacidad
            AND d."Activo" = TRUE
        ORDER BY
            d."IdDocumentoIncapacidadTrabajador" ASC
        '''
    )

    documentos = db.execute(
        consulta_documentos,
        {"id_incapacidad": id_incapacidad},
    ).mappings().all()

    data = _incapacidad_a_dict(fila)
    data["documentos"] = [
        _documento_a_dict(documento)
        for documento in documentos
    ]

    return {
        "success": True,
        "data": data,
    }


@router.put("/{id_incapacidad}/tipo")
def actualizar_tipo_incapacidad_nomina(
    id_incapacidad: int,
    payload: ActualizarTipoIncapacidadRequest,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    tipo_incapacidad = str(payload.tipo_incapacidad or "").strip().upper()
    descripcion_tipo = TIPOS_INCAPACIDAD_NOMINA.get(tipo_incapacidad)

    if not descripcion_tipo:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="El tipo de incapacidad seleccionado no es válido.",
        )

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "TipoIncapacidad" = :tipo_incapacidad,
            "DescripcionTipoIncapacidad" = :descripcion_tipo,
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'REGISTRADA'
        RETURNING
            "IdIncapacidadTrabajador",
            "TipoIncapacidad",
            "DescripcionTipoIncapacidad",
            "Estado",
            "FechaActualizacion"
        """
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
            "tipo_incapacidad": tipo_incapacidad,
            "descripcion_tipo": descripcion_tipo,
        },
    ).mappings().first()

    if not fila:
        db.rollback()

        estado_actual = _consultar_estado_incapacidad(
            db,
            id_incapacidad,
        )

        if not estado_actual or estado_actual["Activo"] is not True:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        estado = _normalizar_estado(
            estado_actual["Estado"]
        )

        if estado.upper() == "BORRADOR":
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "El tipo de incapacidad solo puede modificarse mientras "
                "la incapacidad se encuentre en estado REGISTRADA. "
                f"Estado actual: {estado}."
            ),
        )

    db.commit()

    return {
        "success": True,
        "message": "Tipo de incapacidad actualizado correctamente.",
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "tipo_incapacidad": fila["TipoIncapacidad"],
            "descripcion_tipo": fila["DescripcionTipoIncapacidad"],
            "estado": fila["Estado"],
            "fecha_actualizacion": fila["FechaActualizacion"],
        },
    }


@router.put("/{id_incapacidad}/datos")
def actualizar_datos_incapacidad_nomina(
    id_incapacidad: int,
    payload: ActualizarDatosIncapacidadRequest,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    tipo_incapacidad = str(
        payload.tipo_incapacidad or ""
    ).strip().upper()

    descripcion_tipo = TIPOS_INCAPACIDAD_NOMINA.get(
        tipo_incapacidad
    )

    if not descripcion_tipo:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="El tipo de incapacidad seleccionado no es válido.",
        )

    dias_incapacidad = int(payload.dias_incapacidad)
    fecha_inicio = payload.fecha_inicio
    concepto_sinergy = str(payload.concepto_sinergy or "").strip()
    diagnostico = str(payload.diagnostico or "").strip()

    if len(concepto_sinergy) > 10:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="El concepto permite máximo 10 caracteres.",
        )

    if concepto_sinergy:
        concepto_valido = db.execute(
            text(
                """
                SELECT
                    "CodigoConcepto"
                FROM public."ConceptoIncapacidadSinergy"
                WHERE
                    "CodigoConcepto" = :codigo_concepto
                    AND "Activo" = TRUE
                LIMIT 1
                """
            ),
            {
                "codigo_concepto": concepto_sinergy,
            },
        ).mappings().first()

        if not concepto_valido:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=(
                    f"El concepto {concepto_sinergy} no existe o no está "
                    "activo en el catálogo de conceptos de Sinergy."
                ),
            )

    if len(diagnostico) > 9:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="El diagnóstico permite máximo 9 caracteres.",
        )

    _validar_dias_segun_tipo(
        tipo_incapacidad,
        dias_incapacidad,
    )

    incapacidad_actual = db.execute(
        text(
            """
            SELECT
                "Estado",
                "Activo"
            FROM public."IncapacidadTrabajador"
            WHERE "IdIncapacidadTrabajador" = :id_incapacidad
            LIMIT 1
            """
        ),
        {"id_incapacidad": id_incapacidad},
    ).mappings().first()

    if not incapacidad_actual or incapacidad_actual["Activo"] is not True:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="La incapacidad no existe o no está disponible.",
        )

    estado_actual = _normalizar_estado(
        incapacidad_actual["Estado"]
    )

    if estado_actual.upper() != "REGISTRADA":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "La información de la incapacidad solo puede modificarse "
                "mientras se encuentre en estado REGISTRADA. "
                f"Estado actual: {estado_actual}."
            ),
        )

    fecha_final = fecha_inicio + timedelta(
        days=dias_incapacidad - 1
    )

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "TipoIncapacidad" = :tipo_incapacidad,
            "DescripcionTipoIncapacidad" = :descripcion_tipo,
            "FechaInicio" = :fecha_inicio,
            "DiasIncapacidad" = :dias_incapacidad,
            "FechaFinal" = :fecha_final,
            "EsProrroga" = :es_prorroga,
            "ConceptoSinergy" = :concepto_sinergy,
            "Diagnostico" = :diagnostico,
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'REGISTRADA'
        RETURNING
            "IdIncapacidadTrabajador",
            "TipoIncapacidad",
            "DescripcionTipoIncapacidad",
            "FechaInicio",
            "DiasIncapacidad",
            "FechaFinal",
            "EsProrroga",
            "ConceptoSinergy",
            "Diagnostico",
            "Estado",
            "FechaActualizacion"
        """
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
            "tipo_incapacidad": tipo_incapacidad,
            "descripcion_tipo": descripcion_tipo,
            "fecha_inicio": fecha_inicio,
            "dias_incapacidad": dias_incapacidad,
            "fecha_final": fecha_final,
            "es_prorroga": bool(payload.es_prorroga),
            "concepto_sinergy": concepto_sinergy or None,
            "diagnostico": diagnostico or None,
        },
    ).mappings().first()

    if not fila:
        db.rollback()
        _resolver_incapacidad_no_actualizada(
            db,
            id_incapacidad,
        )

    db.commit()

    return {
        "success": True,
        "message": "Información de la incapacidad actualizada correctamente.",
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "tipo_incapacidad": fila["TipoIncapacidad"],
            "descripcion_tipo": fila["DescripcionTipoIncapacidad"],
            "fecha_inicio": fila["FechaInicio"],
            "dias_incapacidad": fila["DiasIncapacidad"],
            "fecha_final": fila["FechaFinal"],
            "es_prorroga": fila["EsProrroga"],
            "concepto_sinergy": fila["ConceptoSinergy"],
            "diagnostico": fila["Diagnostico"],
            "estado": fila["Estado"],
            "fecha_actualizacion": fila["FechaActualizacion"],
        },
    }


@router.put("/{id_incapacidad}/aprobar")
def aprobar_incapacidad_nomina(
    id_incapacidad: int,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    usuario_gestion = _obtener_usuario_gestion(usuario_actual)

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "Estado" = 'APROBADA',
            "ObservacionNomina" = NULL,
            "UsuarioGestionNomina" = :usuario_gestion,
            "FechaGestionNomina" = CURRENT_TIMESTAMP,
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'REGISTRADA'
        RETURNING
            "IdIncapacidadTrabajador",
            "Estado",
            "ObservacionNomina",
            "UsuarioGestionNomina",
            "FechaGestionNomina",
            "FechaActualizacion"
        """
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
            "usuario_gestion": usuario_gestion,
        },
    ).mappings().first()

    if not fila:
        db.rollback()
        _resolver_incapacidad_no_actualizada(db, id_incapacidad)

    db.commit()

    correo_enviado, detalle_correo = _enviar_notificacion_gestion(
        db=db,
        id_incapacidad=id_incapacidad,
        estado="APROBADA",
    )

    return {
        "success": True,
        "message": "Incapacidad aprobada correctamente.",
        "correo_enviado": correo_enviado,
        "detalle_correo": detalle_correo,
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "estado": fila["Estado"],
            "observacion_nomina": fila["ObservacionNomina"],
            "usuario_gestion_nomina": fila["UsuarioGestionNomina"],
            "fecha_gestion_nomina": fila["FechaGestionNomina"],
            "fecha_actualizacion": fila["FechaActualizacion"],
        },
    }


@router.put("/{id_incapacidad}/rechazar")
def rechazar_incapacidad_nomina(
    id_incapacidad: int,
    payload: RechazoIncapacidadRequest,
    request: Request,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    observacion = payload.observacion.strip()

    if len(observacion) < 3:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Debe registrar un motivo de rechazo válido.",
        )

    usuario_gestion = _obtener_usuario_gestion(usuario_actual)

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "Estado" = 'RECHAZADA',
            "ObservacionNomina" = :observacion,
            "UsuarioGestionNomina" = :usuario_gestion,
            "FechaGestionNomina" = CURRENT_TIMESTAMP,
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'REGISTRADA'
        RETURNING
            "IdIncapacidadTrabajador",
            "Estado",
            "ObservacionNomina",
            "UsuarioGestionNomina",
            "FechaGestionNomina",
            "FechaActualizacion"
        """
    )

    try:
        fila = db.execute(
            consulta,
            {
                "id_incapacidad": id_incapacidad,
                "observacion": observacion,
                "usuario_gestion": usuario_gestion,
            },
        ).mappings().first()

        if not fila:
            db.rollback()
            _resolver_incapacidad_no_actualizada(
                db,
                id_incapacidad,
            )

        correccion, token_plano = _crear_solicitud_correccion(
            db=db,
            id_incapacidad=id_incapacidad,
            motivo_correccion=observacion,
            usuario_solicitud=usuario_gestion,
        )

        url_correccion = _construir_url_correccion(
            request=request,
            token_plano=token_plano,
        )

        db.commit()

    except HTTPException:
        raise

    except Exception as exc:
        db.rollback()

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                "No fue posible registrar el rechazo y generar "
                "la solicitud de corrección."
            ),
        ) from exc

    correo_enviado, detalle_correo = _enviar_notificacion_gestion(
        db=db,
        id_incapacidad=id_incapacidad,
        estado="RECHAZADA",
        observacion=observacion,
        url_correccion=url_correccion,
    )

    return {
        "success": True,
        "message": (
            "Incapacidad rechazada correctamente. "
            "Se generó la solicitud de corrección para el trabajador."
        ),
        "correo_enviado": correo_enviado,
        "detalle_correo": detalle_correo,
        "correccion_generada": True,
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "estado": fila["Estado"],
            "observacion_nomina": fila["ObservacionNomina"],
            "usuario_gestion_nomina": fila["UsuarioGestionNomina"],
            "fecha_gestion_nomina": fila["FechaGestionNomina"],
            "fecha_actualizacion": fila["FechaActualizacion"],
            "id_correccion": correccion["IdCorreccionIncapacidad"],
            "estado_correccion": correccion["EstadoCorreccion"],
            "fecha_solicitud_correccion": correccion["FechaSolicitud"],
            "fecha_expiracion_enlace": correccion[
                "FechaExpiracionToken"
            ],
        },
    }


@router.put("/{id_incapacidad}/pendiente-radicacion")
def marcar_pendiente_radicacion_nomina(
    id_incapacidad: int,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "Estado" = 'PENDIENTE RADICACION',
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'APROBADA'
        RETURNING
            "IdIncapacidadTrabajador",
            "Estado",
            "ObservacionNomina",
            "UsuarioGestionNomina",
            "FechaGestionNomina",
            "FechaActualizacion"
        """
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
        },
    ).mappings().first()

    if not fila:
        db.rollback()

        estado_actual = _consultar_estado_incapacidad(
            db,
            id_incapacidad,
        )

        if not estado_actual or estado_actual["Activo"] is not True:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        estado = _normalizar_estado(
            estado_actual["Estado"]
        )

        if estado.upper() == "BORRADOR":
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Solo una incapacidad en estado APROBADA puede pasar "
                "a PENDIENTE RADICACION. "
                f"Estado actual: {estado}."
            ),
        )

    db.commit()

    return {
        "success": True,
        "message": (
            "Incapacidad marcada como pendiente de radicación correctamente."
        ),
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "estado": fila["Estado"],
            "observacion_nomina": fila["ObservacionNomina"],
            "usuario_gestion_nomina": fila["UsuarioGestionNomina"],
            "fecha_gestion_nomina": fila["FechaGestionNomina"],
            "fecha_actualizacion": fila["FechaActualizacion"],
        },
    }


@router.put("/{id_incapacidad}/radicar")
def radicar_incapacidad_nomina(
    id_incapacidad: int,
    payload: RadicarIncapacidadRequest,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    numero_radicado = payload.numero_radicado.strip()

    if not numero_radicado:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Debe registrar el número de radicado.",
        )

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "Estado" = 'RADICADO',
            "NumeroRadicado" = :numero_radicado,
            "FechaRadicacion" = :fecha_radicacion,
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'PENDIENTE RADICACION'
        RETURNING
            "IdIncapacidadTrabajador",
            "Estado",
            "ObservacionNomina",
            "UsuarioGestionNomina",
            "FechaGestionNomina",
            "NumeroRadicado",
            "FechaRadicacion",
            "FechaActualizacion"
        """
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
            "numero_radicado": numero_radicado,
            "fecha_radicacion": payload.fecha_radicacion,
        },
    ).mappings().first()

    if not fila:
        db.rollback()

        estado_actual = _consultar_estado_incapacidad(
            db,
            id_incapacidad,
        )

        if not estado_actual or estado_actual["Activo"] is not True:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        estado = _normalizar_estado(
            estado_actual["Estado"]
        )

        if estado.upper() == "BORRADOR":
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Solo una incapacidad en estado PENDIENTE RADICACION "
                "puede pasar a RADICADO. "
                f"Estado actual: {estado}."
            ),
        )

    db.commit()

    return {
        "success": True,
        "message": "Incapacidad radicada correctamente.",
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "estado": fila["Estado"],
            "observacion_nomina": fila["ObservacionNomina"],
            "usuario_gestion_nomina": fila["UsuarioGestionNomina"],
            "fecha_gestion_nomina": fila["FechaGestionNomina"],
            "numero_radicado": fila["NumeroRadicado"],
            "fecha_radicacion": fila["FechaRadicacion"],
            "fecha_actualizacion": fila["FechaActualizacion"],
        },
    }


@router.put("/{id_incapacidad}/negar")
def negar_incapacidad_nomina(
    id_incapacidad: int,
    payload: NegarIncapacidadRequest,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    causal_negacion = payload.causal_negacion.strip()

    if len(causal_negacion) < 3:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Debe registrar una causal de negación válida.",
        )

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "Estado" = 'NEGADO',
            "CausalNegacion" = :causal_negacion,
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'RADICADO'
        RETURNING
            "IdIncapacidadTrabajador",
            "Estado",
            "ObservacionNomina",
            "UsuarioGestionNomina",
            "FechaGestionNomina",
            "NumeroRadicado",
            "FechaRadicacion",
            "CausalNegacion",
            "FechaActualizacion"
        """
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
            "causal_negacion": causal_negacion,
        },
    ).mappings().first()

    if not fila:
        db.rollback()

        estado_actual = _consultar_estado_incapacidad(
            db,
            id_incapacidad,
        )

        if not estado_actual or estado_actual["Activo"] is not True:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        estado = _normalizar_estado(
            estado_actual["Estado"]
        )

        if estado.upper() == "BORRADOR":
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Solo una incapacidad en estado RADICADO puede pasar "
                "a NEGADO. "
                f"Estado actual: {estado}."
            ),
        )

    db.commit()

    correo_enviado, detalle_correo = _enviar_notificacion_gestion(
        db=db,
        id_incapacidad=id_incapacidad,
        estado="NEGADO",
        observacion=causal_negacion,
    )

    return {
        "success": True,
        "message": "Incapacidad marcada como negada correctamente.",
        "correo_enviado": correo_enviado,
        "detalle_correo": detalle_correo,
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "estado": fila["Estado"],
            "observacion_nomina": fila["ObservacionNomina"],
            "usuario_gestion_nomina": fila["UsuarioGestionNomina"],
            "fecha_gestion_nomina": fila["FechaGestionNomina"],
            "numero_radicado": fila["NumeroRadicado"],
            "fecha_radicacion": fila["FechaRadicacion"],
            "causal_negacion": fila["CausalNegacion"],
            "fecha_actualizacion": fila["FechaActualizacion"],
        },
    }


@router.put("/{id_incapacidad}/en-proceso-pago")
def marcar_en_proceso_pago_nomina(
    id_incapacidad: int,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "Estado" = 'EN PROCESO DE PAGO',
            "CausalNegacion" = NULL,
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'RADICADO'
        RETURNING
            "IdIncapacidadTrabajador",
            "Estado",
            "ObservacionNomina",
            "UsuarioGestionNomina",
            "FechaGestionNomina",
            "NumeroRadicado",
            "FechaRadicacion",
            "CausalNegacion",
            "FechaActualizacion"
        """
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
        },
    ).mappings().first()

    if not fila:
        db.rollback()

        estado_actual = _consultar_estado_incapacidad(
            db,
            id_incapacidad,
        )

        if not estado_actual or estado_actual["Activo"] is not True:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        estado = _normalizar_estado(
            estado_actual["Estado"]
        )

        if estado.upper() == "BORRADOR":
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Solo una incapacidad en estado RADICADO puede pasar "
                "a EN PROCESO DE PAGO. "
                f"Estado actual: {estado}."
            ),
        )

    db.commit()

    correo_enviado, detalle_correo = _enviar_notificacion_gestion(
        db=db,
        id_incapacidad=id_incapacidad,
        estado="EN PROCESO DE PAGO",
    )

    return {
        "success": True,
        "message": (
            "Incapacidad marcada como en proceso de pago correctamente."
        ),
        "correo_enviado": correo_enviado,
        "detalle_correo": detalle_correo,
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "estado": fila["Estado"],
            "observacion_nomina": fila["ObservacionNomina"],
            "usuario_gestion_nomina": fila["UsuarioGestionNomina"],
            "fecha_gestion_nomina": fila["FechaGestionNomina"],
            "numero_radicado": fila["NumeroRadicado"],
            "fecha_radicacion": fila["FechaRadicacion"],
            "causal_negacion": fila["CausalNegacion"],
            "fecha_actualizacion": fila["FechaActualizacion"],
        },
    }


@router.put("/{id_incapacidad}/pagar")
def pagar_incapacidad_nomina(
    id_incapacidad: int,
    payload: PagarIncapacidadRequest,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    valor_pagado = payload.valor_pagado.quantize(
        Decimal("0.01")
    )

    consulta = text(
        """
        UPDATE public."IncapacidadTrabajador"
        SET
            "Estado" = 'PAGADO',
            "ValorPagado" = :valor_pagado,
            "FechaActualizacion" = CURRENT_TIMESTAMP
        WHERE
            "IdIncapacidadTrabajador" = :id_incapacidad
            AND "Activo" = TRUE
            AND UPPER(COALESCE("Estado", '')) = 'EN PROCESO DE PAGO'
        RETURNING
            "IdIncapacidadTrabajador",
            "Estado",
            "ObservacionNomina",
            "UsuarioGestionNomina",
            "FechaGestionNomina",
            "NumeroRadicado",
            "FechaRadicacion",
            "CausalNegacion",
            "ValorPagado",
            "FechaActualizacion"
        """
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
            "valor_pagado": valor_pagado,
        },
    ).mappings().first()

    if not fila:
        db.rollback()

        estado_actual = _consultar_estado_incapacidad(
            db,
            id_incapacidad,
        )

        if not estado_actual or estado_actual["Activo"] is not True:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        estado = _normalizar_estado(
            estado_actual["Estado"]
        )

        if estado.upper() == "BORRADOR":
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="La incapacidad no existe o no está disponible.",
            )

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Solo una incapacidad en estado EN PROCESO DE PAGO "
                "puede pasar a PAGADO. "
                f"Estado actual: {estado}."
            ),
        )

    db.commit()

    correo_enviado, detalle_correo = _enviar_notificacion_gestion(
        db=db,
        id_incapacidad=id_incapacidad,
        estado="PAGADO",
    )

    return {
        "success": True,
        "message": "Pago de incapacidad registrado correctamente.",
        "correo_enviado": correo_enviado,
        "detalle_correo": detalle_correo,
        "data": {
            "id_incapacidad": fila["IdIncapacidadTrabajador"],
            "estado": fila["Estado"],
            "observacion_nomina": fila["ObservacionNomina"],
            "usuario_gestion_nomina": fila["UsuarioGestionNomina"],
            "fecha_gestion_nomina": fila["FechaGestionNomina"],
            "numero_radicado": fila["NumeroRadicado"],
            "fecha_radicacion": fila["FechaRadicacion"],
            "causal_negacion": fila["CausalNegacion"],
            "valor_pagado": fila["ValorPagado"],
            "fecha_actualizacion": fila["FechaActualizacion"],
        },
    }


@router.get("/{id_incapacidad}/documentos-consolidados")
def descargar_documentos_consolidados_nomina(
    id_incapacidad: int,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    consulta = text(
        """
        SELECT
            d."IdDocumentoIncapacidadTrabajador",
            d."IdIncapacidadTrabajador",
            d."TipoDocumento",
            d."NombreArchivo",
            d."Formato",
            d."MimeType",
            d."DocumentoCargado",
            rp."NumeroIdentificacion",
            i."FechaInicio"
        FROM public."DocumentoIncapacidadTrabajador" d
        INNER JOIN public."IncapacidadTrabajador" i
            ON i."IdIncapacidadTrabajador" = d."IdIncapacidadTrabajador"
        INNER JOIN public."RegistroPersonal" rp
            ON rp."IdRegistroPersonal" = i."IdRegistroPersonal"
        WHERE
            i."IdIncapacidadTrabajador" = :id_incapacidad
            AND i."Activo" = TRUE
            AND d."Activo" = TRUE
            AND COALESCE(i."Estado", '') <> 'BORRADOR'
        ORDER BY
            d."IdDocumentoIncapacidadTrabajador" ASC
        """
    )

    documentos = db.execute(
        consulta,
        {"id_incapacidad": id_incapacidad},
    ).mappings().all()

    if not documentos:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="La incapacidad no tiene documentos activos para descargar.",
        )

    pdf_consolidado = _construir_pdf_consolidado(documentos)

    identificacion = (
        str(documentos[0]["NumeroIdentificacion"] or "")
        .strip()
        .replace(" ", "_")
    )

    fecha_inicio = documentos[0]["FechaInicio"]

    if fecha_inicio:
        try:
            fecha_texto = fecha_inicio.strftime("%Y-%m-%d")
        except AttributeError:
            fecha_texto = str(fecha_inicio).replace("/", "-")
    else:
        fecha_texto = "sin_fecha"

    nombre_archivo = (
        f"incapacidad_{identificacion or id_incapacidad}_"
        f"{fecha_texto}_soportes.pdf"
    )

    return Response(
        content=pdf_consolidado,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{nombre_archivo}"'
            ),
            "Cache-Control": "no-store",
        },
    )


@router.get("/{id_incapacidad}/documentos/{id_documento}")
def ver_documento_incapacidad_nomina(
    id_incapacidad: int,
    id_documento: int,
    db: Session = Depends(get_db),
    usuario_actual=Depends(get_current_user),
):
    _validar_acceso_nomina(usuario_actual)

    consulta = text(
        '''
        SELECT
            d."IdDocumentoIncapacidadTrabajador",
            d."IdIncapacidadTrabajador",
            d."NombreArchivo",
            d."MimeType",
            d."DocumentoCargado"
        FROM public."DocumentoIncapacidadTrabajador" d
        INNER JOIN public."IncapacidadTrabajador" i
            ON i."IdIncapacidadTrabajador" = d."IdIncapacidadTrabajador"
        WHERE
            i."IdIncapacidadTrabajador" = :id_incapacidad
            AND d."IdDocumentoIncapacidadTrabajador" = :id_documento
            AND i."Activo" = TRUE
            AND d."Activo" = TRUE
            AND COALESCE(i."Estado", '') <> 'BORRADOR'
        LIMIT 1
        '''
    )

    fila = db.execute(
        consulta,
        {
            "id_incapacidad": id_incapacidad,
            "id_documento": id_documento,
        },
    ).mappings().first()

    if not fila:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="El documento no existe o no está disponible.",
        )

    contenido = fila["DocumentoCargado"]

    if not contenido:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="El documento no contiene información para visualizar.",
        )

    nombre_archivo = (
        str(fila["NombreArchivo"] or "documento_incapacidad")
        .replace('"', "")
        .replace("\r", "")
        .replace("\n", "")
    )

    mime_type = (
        str(fila["MimeType"] or "").strip()
        or "application/octet-stream"
    )

    return Response(
        content=bytes(contenido),
        media_type=mime_type,
        headers={
            "Content-Disposition": (
                f'inline; filename="{nombre_archivo}"'
            ),
            "Cache-Control": "no-store",
        },
    )
